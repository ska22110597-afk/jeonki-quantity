"""내역서 읽기(원본 미수정) 및 3시트 결과 엑셀 생성."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.estimate_parse import (
    EstimateSheet,
    find_column_index,
    find_header_row,
    find_quantity_column,
    first_data_row_number,
    is_merge_top_left,
    is_external_formula,
    is_section_row,
    is_sundry_form_row,
    load_estimate_sheet,
    read_named_sheet_rows,
    read_workbook_first_sheet,
)
from app.paths import assert_safe_save, build_result_path
from app.pumsam import (
    PUMSAM_SHEET_NAME,
    PumsamRow,
    ensure_all_pumsam_databases,
    import_pumsam_file,
    load_pumsam_database,
    merge_pumsam_rows,
    pumsam_surcharge_note,
    rows_from_grid,
    save_pumsam_database,
)

ESTIMATE_SHEET_NAME = "내역서"
ILWIDAE_LIST_SHEET_NAME = "일위대가목록"
QUANTITY_SHEET_NAME = "공량산출서"
COMPARE_SHEET_NAME = "단가대비표"
ILWIDAE_SHEET_NAME = "일위대가"
UNIT_PRICE_SHEET_NAME = COMPARE_SHEET_NAME

ROW_HEIGHT = 20
FORM_ROW_HEIGHT = 30
PUMSAM_DATA_START = 5
QTY_HEADER_ROW = 2
QTY_SUBHEADER_ROW = 3
COMPARE_DATA_START = 5
COMPARE_PRICE_COL = 11  # L열 적용단가. 코드 열은 쓰지 않는다.

WHITE = PatternFill("solid", fgColor="FFFFFF")
TITLE_FONT = Font(name="굴림", size=16, bold=True, underline="single", color="000000")
HEADER_FONT = Font(name="굴림", size=11, bold=True, color="000000")
BODY_FONT = Font(name="굴림", size=11, bold=False, color="000000")
SECTION_FONT = Font(name="굴림", size=11, bold=False, color="000000")
THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=False)

NUMBER_FORMAT = "#,##0.000"
QTY_FORMAT = "#,##0"
PRICE_FORMAT = "#,##0.00"
AMOUNT_FORMAT = "#,##0.0"
PAGE_FORMAT = "0"
MONEY_FORMAT = PRICE_FORMAT
PERCENT_FORMAT = "0%"
RATE_FORMAT = "0"
PUMSAM_FORMAT = "0.000"

QTY_LAST_COL = 12


def concat_formula(row: int) -> str:
    return f"=CONCATENATE(B{row},C{row})"


def decided_qty_formula(row: int) -> str:
    """결정수량 = TRUNC(산출수량 × (1+할증), 0)."""
    return f"=TRUNC(G{row}*(1+F{row}),0)"


def source_qty_formula(source_col_letter: str, row: int, sheet_name: str | None = None) -> str:
    """산출수량 = 같은 행의 일위대가목록 수량."""
    name = sheet_name or ILWIDAE_LIST_SHEET_NAME
    return f"='{name}'!{source_col_letter}{row}"


def gongryang_formula(row: int, last_row: int | None = None) -> str:
    """공량. 품셈표에 인부가 여러 명이면 품셈×할증을 모두 더한다."""
    if last_row is None:
        return f'=IF(G{row}*I{row}=0,"",G{row}*I{row}*(J{row}/100))'
    start = PUMSAM_DATA_START
    return (
        f'=IF(G{row}=0,"",G{row}*SUMPRODUCT('
        f"('{PUMSAM_SHEET_NAME}'!$B${start}:$B${last_row}=B{row})*"
        f"('{PUMSAM_SHEET_NAME}'!$C${start}:$C${last_row}=C{row})*"
        f"('{PUMSAM_SHEET_NAME}'!$F${start}:$F${last_row})*"
        f"('{PUMSAM_SHEET_NAME}'!$G${start}:$G${last_row}/100)))"
    )


def pumsam_vlookup(row: int, col: int, last_row: int) -> str:
    return (
        f'=IFERROR(VLOOKUP($A{row},'
        f"'{PUMSAM_SHEET_NAME}'!$A${PUMSAM_DATA_START}:$H${last_row},{col},FALSE),\"\")"
    )


def labor_formula(row: int, last_row: int) -> str:
    return pumsam_vlookup(row, 5, last_row)


def pumsam_formula(row: int, last_row: int) -> str:
    return pumsam_vlookup(row, 6, last_row)


def labor_rate_formula(row: int, last_row: int) -> str:
    return pumsam_vlookup(row, 7, last_row)


def ref_formula(row: int, last_row: int) -> str:
    return pumsam_vlookup(row, 8, last_row)


# 이전 이름 호환
def calc_qty_formula(row: int) -> str:
    return decided_qty_formula(row)


def gongryang_sum_formula(end_row: int) -> str:
    return f"=SUM(K{max(first_qty_data_row_fallback(), 4)}:K{end_row})"


def first_qty_data_row_fallback() -> int:
    return 5


def _unmerge_all(sheet: Worksheet) -> None:
    for rng in list(sheet.merged_cells.ranges):
        try:
            sheet.unmerge_cells(str(rng))
        except ValueError:
            continue


def _ranges_overlap(
    left: tuple[int, int, int, int],
    right: tuple[int, int, int, int],
) -> bool:
    a_min_r, a_min_c, a_max_r, a_max_c = left
    b_min_r, b_min_c, b_max_r, b_max_c = right
    return not (
        a_max_r < b_min_r or a_min_r > b_max_r or a_max_c < b_min_c or a_min_c > b_max_c
    )


def _without_overlapping_merges(
    merges: list[tuple[int, int, int, int]],
) -> list[tuple[int, int, int, int]]:
    """겹치는 병합은 엑셀 복구 대화상자를 만든다. 앞선 범위만 남긴다."""
    kept: list[tuple[int, int, int, int]] = []
    for item in merges:
        min_row, min_col, max_row, max_col = item
        if min_row == max_row and min_col == max_col:
            continue
        if any(_ranges_overlap(item, prev) for prev in kept):
            continue
        kept.append(item)
    return kept


def _safe_merge(
    sheet: Worksheet,
    min_row: int,
    min_col: int,
    max_row: int,
    max_col: int,
) -> None:
    if min_row < 1 or min_col < 1 or max_row < min_row or max_col < min_col:
        return
    if min_row == max_row and min_col == max_col:
        return
    new = (min_row, min_col, max_row, max_col)
    for rng in list(sheet.merged_cells.ranges):
        existing = (int(rng.min_row), int(rng.min_col), int(rng.max_row), int(rng.max_col))
        if _ranges_overlap(new, existing):
            try:
                sheet.unmerge_cells(str(rng))
            except ValueError:
                continue
    try:
        sheet.merge_cells(
            start_row=min_row,
            start_column=min_col,
            end_row=max_row,
            end_column=max_col,
        )
    except ValueError:
        return


def write_title_banner(
    sheet: Worksheet,
    title: str,
    last_col: int,
    *,
    row_height: int = FORM_ROW_HEIGHT,
) -> None:
    """1~2행을 병합한 가운데 제목. 밑줄."""
    for col in range(1, last_col + 1):
        _set_cell(sheet, 1, col, title if col == 1 else None, font=TITLE_FONT, align=CENTER)
        _set_cell(sheet, 2, col, None, font=TITLE_FONT, align=CENTER)
    _safe_merge(sheet, 1, 1, 2, last_col)
    sheet.row_dimensions[1].height = row_height
    sheet.row_dimensions[2].height = row_height


def _apply_sheet_look(
    sheet: Worksheet,
    max_row: int,
    max_col: int,
    *,
    row_height: int | None = None,
) -> None:
    height = ROW_HEIGHT if row_height is None else row_height
    sheet.sheet_properties.tabColor = "FFFFFF"
    sheet.sheet_format.defaultRowHeight = height
    for r in range(1, max(max_row, 1) + 1):
        sheet.row_dimensions[r].height = height
        for c in range(1, max(max_col, 1) + 1):
            cell = sheet.cell(row=r, column=c)
            cell.fill = WHITE
            align = cell.alignment
            cell.alignment = Alignment(
                horizontal=align.horizontal,
                vertical="center",
                wrap_text=False,
            )


def _set_cell(
    sheet: Worksheet,
    row: int,
    col: int,
    value: Any,
    *,
    font: Font | None = None,
    align: Alignment | None = None,
    number_format: str | None = None,
    border: bool = True,
) -> None:
    cell = sheet.cell(row=row, column=col, value=value)
    cell.font = font or BODY_FONT
    cell.fill = WHITE
    cell.alignment = align or LEFT
    if border:
        cell.border = THIN
    if number_format:
        cell.number_format = number_format


def _write_estimate_sheet(sheet: Worksheet, estimate: EstimateSheet) -> None:
    if estimate.title:
        sheet.title = estimate.title
    filled = estimate.filled
    raw = estimate.raw
    max_row = estimate.max_row
    max_col = min(max(estimate.max_col, 13), 20)
    header_idx = find_header_row(filled)
    name_idx = find_column_index(filled[header_idx], "명칭") if filled else None
    spec_idx = find_column_index(filled[header_idx], "규격") if filled else None
    unit_idx = find_column_index(filled[header_idx], "단위") if filled else None
    data_start = first_data_row_number(filled) if filled else 2
    subheader = filled[header_idx + 1] if filled and header_idx + 1 < len(filled) else []
    _unmerge_all(sheet)

    clipped: list[tuple[int, int, int, int]] = []
    for min_row, min_col, max_r, max_c in estimate.merges:
        if min_row < 1 or min_col < 1:
            continue
        end_row = min(max_r, max_row)
        end_col = min(max_c, max_col)
        if end_row < min_row or end_col < min_col:
            continue
        clipped.append((min_row, min_col, end_row, end_col))
    kept_merges = _without_overlapping_merges(clipped)

    def _col_number_format(col: int) -> str:
        token = ""
        if col - 1 < len(subheader):
            token = str(subheader[col - 1] or "")
        if "금액" in token.replace(" ", ""):
            return AMOUNT_FORMAT
        return PRICE_FORMAT

    for r_idx in range(1, max_row + 1):
        source = raw[r_idx - 1] if r_idx - 1 < len(raw) else []
        filled_row = filled[r_idx - 1] if r_idx - 1 < len(filled) else []
        name = filled_row[name_idx] if name_idx is not None and name_idx < len(filled_row) else None
        spec = filled_row[spec_idx] if spec_idx is not None and spec_idx < len(filled_row) else None
        unit = filled_row[unit_idx] if unit_idx is not None and unit_idx < len(filled_row) else None
        section = r_idx >= data_start and is_section_row(name, spec, unit) and name is not None
        is_header = r_idx <= data_start - 1

        for c_idx in range(1, max_col + 1):
            if not is_merge_top_left(r_idx, c_idx, kept_merges):
                cell = sheet.cell(row=r_idx, column=c_idx)
                cell.value = None
                cell.fill = WHITE
                cell.border = THIN
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
                cell.font = HEADER_FONT if is_header else BODY_FONT
                continue
            value = source[c_idx - 1] if c_idx - 1 < len(source) else None
            if value in (None, "") and c_idx - 1 < len(filled_row):
                filled_value = filled_row[c_idx - 1]
                if isinstance(filled_value, str) and filled_value.startswith("="):
                    value = filled_value
            if is_external_formula(value):
                value = None
            font = HEADER_FONT if is_header else (SECTION_FONT if section and c_idx == 1 else BODY_FONT)
            align = CENTER if is_header or c_idx in (3, 4) else LEFT
            number_format = None
            if r_idx >= data_start and isinstance(value, (int, float)):
                number_format = _col_number_format(c_idx)
                align = RIGHT
            elif isinstance(value, str) and value.startswith("=") and c_idx >= 4:
                number_format = _col_number_format(c_idx)
                align = RIGHT
            _set_cell(
                sheet,
                r_idx,
                c_idx,
                value,
                font=font,
                align=align,
                number_format=number_format,
            )

    for min_row, min_col, end_row, end_col in kept_merges:
        _safe_merge(sheet, min_row, min_col, end_row, end_col)

    _apply_sheet_look(sheet, max_row, max_col, row_height=FORM_ROW_HEIGHT)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 8
    sheet.column_dimensions["D"].width = 10
    for col in range(5, max_col + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 12


def _write_pumsam_sheet(sheet: Worksheet, rows: list[PumsamRow]) -> None:
    last_col = 9
    write_title_banner(sheet, "품 셈 표", last_col)
    header_row = 3
    sub_row = 4
    _set_cell(sheet, header_row, 1, "품목", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, header_row, 2, "명칭", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, header_row, 3, "규격", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, header_row, 4, "단위", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, header_row, 5, "공량산출", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, header_row, 9, "비고", font=HEADER_FONT, align=CENTER)
    for col in (6, 7, 8):
        _set_cell(sheet, header_row, col, None, font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 5, "명칭", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 6, "품셈", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 7, "할증%", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 8, "품셈근거", font=HEADER_FONT, align=CENTER)
    for col in (1, 2, 3, 4, 9):
        _set_cell(sheet, sub_row, col, None, font=HEADER_FONT, align=CENTER)

    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    sheet.merge_cells("D3:D4")
    sheet.merge_cells("E3:H3")
    sheet.merge_cells("I3:I4")

    last_data = PUMSAM_DATA_START - 1
    for offset, row in enumerate(rows):
        excel_row = PUMSAM_DATA_START + offset
        last_data = excel_row
        values = [
            concat_formula(excel_row),
            row.get("명칭"),
            row.get("규격"),
            row.get("단위"),
            row.get("노무명칭"),
            row.get("품셈"),
            row.get("할증%"),
            row.get("품셈근거"),
            pumsam_surcharge_note(row),
        ]
        for c_idx, value in enumerate(values, start=1):
            align = CENTER if c_idx in (3, 4, 7) else LEFT
            number_format = None
            if c_idx == 6:
                number_format = PUMSAM_FORMAT
                align = RIGHT
            elif c_idx == 7:
                number_format = RATE_FORMAT
                align = RIGHT
            _set_cell(
                sheet,
                excel_row,
                c_idx,
                value,
                font=BODY_FONT,
                align=align,
                number_format=number_format,
            )

    max_row = max(last_data, 4)
    _apply_sheet_look(sheet, max_row, last_col)
    sheet.row_dimensions[1].height = FORM_ROW_HEIGHT
    sheet.row_dimensions[2].height = FORM_ROW_HEIGHT
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 8
    sheet.column_dimensions["E"].width = 12
    sheet.column_dimensions["F"].width = 10
    sheet.column_dimensions["G"].width = 10
    sheet.column_dimensions["H"].width = 12
    sheet.column_dimensions["I"].width = 28


def _pick(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _write_quantity_header(sheet: Worksheet, header_row: int = 2) -> None:
    sub_row = header_row + 1
    headers = {
        1: "품목",
        2: "명칭",
        3: "규격",
        4: "단위",
        5: "수량",
        8: "공량산출",
        12: "품셈근거",
    }
    for col in range(1, QTY_LAST_COL + 1):
        _set_cell(sheet, header_row, col, headers.get(col), font=HEADER_FONT, align=CENTER)
        _set_cell(sheet, sub_row, col, None, font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 5, "결정수량", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 6, "할증", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 7, "산출수량", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 8, "명칭", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 9, "품셈", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 10, "할증%", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, sub_row, 11, "공량", font=HEADER_FONT, align=CENTER)
    _safe_merge(sheet, header_row, 1, sub_row, 1)
    _safe_merge(sheet, header_row, 2, sub_row, 2)
    _safe_merge(sheet, header_row, 3, sub_row, 3)
    _safe_merge(sheet, header_row, 4, sub_row, 4)
    _safe_merge(sheet, header_row, 5, header_row, 7)
    _safe_merge(sheet, header_row, 8, header_row, 11)
    _safe_merge(sheet, header_row, 12, sub_row, 12)


def _write_quantity_sheet(
    sheet: Worksheet,
    estimate: EstimateSheet,
    pumsam_last_row: int,
    pumsam_rows: list[PumsamRow] | None = None,
    source_sheet_name: str | None = None,
) -> None:
    _unmerge_all(sheet)
    filled = estimate.filled
    header_idx = find_header_row(filled) if filled else 0
    header = filled[header_idx] if filled else []
    name_idx = find_column_index(header, "명칭")
    spec_idx = find_column_index(header, "규격")
    unit_idx = find_column_index(header, "단위")
    qty_idx = find_quantity_column(header)
    qty_letter = get_column_letter(qty_idx + 1) if qty_idx is not None else "D"
    data_start = first_data_row_number(filled) if filled else 4
    qty_source = source_sheet_name or estimate.title or ILWIDAE_LIST_SHEET_NAME

    if data_start >= 5:
        write_title_banner(sheet, "공 량 산 출 서", QTY_LAST_COL)
        _write_quantity_header(sheet, 3)
    elif data_start >= 4:
        for col in range(1, QTY_LAST_COL + 1):
            _set_cell(
                sheet,
                1,
                col,
                "공 량 산 출 서" if col == 1 else None,
                font=TITLE_FONT,
                align=CENTER,
            )
        _safe_merge(sheet, 1, 1, 1, QTY_LAST_COL)
        _write_quantity_header(sheet, 2)
    else:
        labels = ["", "명칭", "규격", "단위", "결정수량", "할증", "산출수량", "명칭", "품셈", "할증%", "공량", "품셈근거"]
        for col, label in enumerate(labels, start=1):
            _set_cell(sheet, 1, col, label or None, font=HEADER_FONT, align=CENTER)

    last_row = max(data_start, len(filled), 3)
    for excel_row in range(data_start, len(filled) + 1):
        source = filled[excel_row - 1]
        name = _pick(source, name_idx)
        spec = _pick(source, spec_idx)
        unit = _pick(source, unit_idx)
        last_row = excel_row
        if name is None and spec is None:
            for col in range(1, QTY_LAST_COL + 1):
                _set_cell(sheet, excel_row, col, None)
            continue

        form_row = is_sundry_form_row(name)
        section = is_section_row(name, spec, unit) and not form_row
        _set_cell(sheet, excel_row, 1, concat_formula(excel_row), font=SECTION_FONT if section else BODY_FONT)
        _set_cell(
            sheet,
            excel_row,
            2,
            name,
            font=SECTION_FONT if section else BODY_FONT,
            align=LEFT,
        )
        _set_cell(sheet, excel_row, 3, None if section else spec, font=BODY_FONT, align=LEFT)
        _set_cell(sheet, excel_row, 4, None if section else unit, font=BODY_FONT, align=CENTER)

        if section or form_row:
            for col in range(5, QTY_LAST_COL + 1):
                _set_cell(sheet, excel_row, col, None)
            continue

        _set_cell(
            sheet,
            excel_row,
            5,
            decided_qty_formula(excel_row),
            align=RIGHT,
            number_format=QTY_FORMAT,
        )
        _set_cell(sheet, excel_row, 6, "=0", align=RIGHT, number_format=PERCENT_FORMAT)
        _set_cell(
            sheet,
            excel_row,
            7,
            source_qty_formula(qty_letter, excel_row, qty_source),
            align=RIGHT,
            number_format=QTY_FORMAT,
        )
        lookup_last = max(pumsam_last_row, PUMSAM_DATA_START)
        _set_cell(sheet, excel_row, 8, labor_formula(excel_row, lookup_last), align=LEFT)
        _set_cell(
            sheet,
            excel_row,
            9,
            pumsam_formula(excel_row, lookup_last),
            align=RIGHT,
            number_format=PUMSAM_FORMAT,
        )
        _set_cell(
            sheet,
            excel_row,
            10,
            labor_rate_formula(excel_row, lookup_last),
            align=RIGHT,
            number_format=RATE_FORMAT,
        )
        _set_cell(
            sheet,
            excel_row,
            11,
            gongryang_formula(excel_row, lookup_last),
            align=RIGHT,
            number_format=NUMBER_FORMAT,
        )
        _set_cell(sheet, excel_row, 12, ref_formula(excel_row, lookup_last), align=LEFT)

    _apply_sheet_look(sheet, last_row, QTY_LAST_COL)
    if data_start >= 5:
        sheet.row_dimensions[1].height = FORM_ROW_HEIGHT
        sheet.row_dimensions[2].height = FORM_ROW_HEIGHT
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 32
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 8
    for col, width in enumerate([12, 10, 12, 12, 10, 10, 12, 12], start=5):
        sheet.column_dimensions[get_column_letter(col)].width = width


def collect_pumsam_rows(
    source_path: Path | None,
    dest_dir: Path | None,
    extra_pumsam_path: Path | None = None,
    db_dir: Path | None = None,
    discipline: str | None = None,
) -> list[PumsamRow]:
    database_dir = db_dir if db_dir is not None else dest_dir
    ensure_all_pumsam_databases(database_dir)
    groups = [load_pumsam_database(database_dir, discipline=discipline)]
    if source_path is not None:
        embedded = read_named_sheet_rows(source_path, "품셈")
        if embedded:
            groups.append(rows_from_grid(embedded))
    if extra_pumsam_path is not None:
        groups.append(import_pumsam_file(extra_pumsam_path))
    merged = merge_pumsam_rows(*groups)
    save_pumsam_database(merged, database_dir, discipline=discipline)
    return merged


def create_result_workbook(
    estimate: EstimateSheet | list[list[Any]],
    pumsam_rows: list[PumsamRow] | None = None,
    source_name: str = "",  # noqa: ARG001 — 이전 호출 호환
) -> Workbook:
    if isinstance(estimate, list):
        estimate = EstimateSheet(filled=estimate, raw=estimate, merges=[])
    rows = pumsam_rows or []
    pumsam_last = PUMSAM_DATA_START + len(rows) - 1 if rows else PUMSAM_DATA_START

    workbook = Workbook()
    sheet1 = workbook.active
    sheet1.title = estimate.title or ILWIDAE_LIST_SHEET_NAME
    _write_estimate_sheet(sheet1, estimate)

    sheet2 = workbook.create_sheet(PUMSAM_SHEET_NAME)
    _write_pumsam_sheet(sheet2, rows)

    sheet3 = workbook.create_sheet(QUANTITY_SHEET_NAME)
    _write_quantity_sheet(
        sheet3,
        estimate,
        pumsam_last,
        rows,
        source_sheet_name=estimate.title or ILWIDAE_LIST_SHEET_NAME,
    )
    return workbook


def save_result_workbook(
    source_path: Path | None = None,
    dest_dir: Path | None = None,
    estimate_rows: list[list[Any]] | None = None,
    extra_pumsam_path: Path | None = None,
    db_dir: Path | None = None,
    unit_price_path: Path | None = None,
    ilwidae_path: Path | None = None,
    estimate_path: Path | None = None,
    discipline: str | None = None,
    mode: str | None = None,
) -> Path:
    """드롭한 원본은 읽기만 하고, 결과 엑셀만 새로 저장한다."""
    from app.pipeline import run_pipeline

    if estimate_path is None and source_path is not None:
        estimate_path = source_path
    return run_pipeline(
        dest_dir=dest_dir,
        unit_price_path=unit_price_path,
        ilwidae_path=ilwidae_path,
        estimate_path=estimate_path,
        extra_pumsam_path=extra_pumsam_path,
        db_dir=db_dir,
        estimate_rows=estimate_rows,
        discipline=discipline,
        mode=mode,
    )


# 이전 이름 호환
read_unit_price_table = read_workbook_first_sheet
