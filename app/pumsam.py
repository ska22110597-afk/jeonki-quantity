"""품셈표 데이터베이스. C드라이브 로컬 파일에 쌓고, 결과 엑셀에도 시트로 넣는다."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.estimate_parse import (
    display_keyword,
    find_column_index,
    find_header_row,
    header_row_span,
    lookup_key,
    normalize_header,
)
from app.pumsam_aliases import PLACE_SURCHARGE_RATES, alias_names, split_place_name
from app.pumsam_text import display_spec, lookup_measure, pumsam_ref_sort_key
from app.merge_parse import SheetRows, fill_merged_values, trim_grid
from app.discipline import (
    ELECTRIC,
    LEGACY_PUMSAM_FILENAME,
    TELECOM,
    normalize_discipline,
    pumsam_filename,
)
from app.electric_pumsam_data import ELECTRIC_RULES, ELECTRIC_TRADE_GUIDE, electric_pumsam_rows
from app.telecom_pumsam_data import TELECOM_RULES, TELECOM_TRADE_GUIDE, telecom_pumsam_rows
from app.paths import (
    bundled_data_dir,
    ensure_result_directory,
    get_result_directory,
    is_allowed_excel,
    user_database_dir,
)

PUMSAM_SHEET_NAME = "품셈표"
PUMSAM_DB_FILENAME = LEGACY_PUMSAM_FILENAME
LEGACY_PUMSAM_DB_FILENAME = "품셈표_데이터베이스.xlsx"
PUMSAM_HEADERS = ["검색키", "명칭", "규격", "단위", "노무명칭", "품셈", "할증%", "품셈근거"]
PUMSAM_DISPLAY_HEADERS = ["키워드", "명칭", "규격", "단위", "노무명칭", "품셈", "할증%", "품셈근거"]
PUMSAM_COLUMN_WIDTHS = (60, 35, 45, 9.38, 15, 10, 5, 15)
PUMSAM_DATA_ALIGNS = ("left", "left", "left", "center", "center", "right", "center", "center")
PUMSAM_ROW_HEIGHT = 20
GULIM = Font(name="굴림", size=11)
GULIM_HEADER = Font(name="굴림", size=11, bold=True)
SKY_BLUE = PatternFill("solid", fgColor="B7DEE8")
WHITE = PatternFill("solid", fgColor="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)

PumsamRow = dict[str, Any]


def format_pumsam_ref(value: Any) -> str:
    """전기5-1 → 전기 5-1. 비고란에 그대로 넣는다."""
    text = str(value or "").strip()
    if not text:
        return ""
    compact = text.replace(" ", "")
    matched = re.fullmatch(r"(전기|통신)(\d+)-(\d+)", compact)
    if matched:
        return f"{matched.group(1)} {matched.group(2)}-{matched.group(3)}"
    return text


def pumsam_name_keys(row: PumsamRow) -> set[str]:
    keys: set[str] = set()
    for field in ("명칭", "짧은명칭"):
        token = lookup_key(row.get(field), "")
        if token:
            keys.add(token)
    return keys


def pumsam_qty_value(row: PumsamRow) -> float:
    value = row.get("품셈")
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return 0.0


def has_pumsam_qty(row: PumsamRow) -> bool:
    """원표에 품 숫자가 있는 인부만 일위대가에 넣는다. '-'·빈칸은 그 직종을 쓰지 않는다는 뜻."""
    value = row.get("품셈")
    if value is None or value == "":
        return False
    if isinstance(value, str) and value.strip() in {"-", "－", "—", "–"}:
        return False
    try:
        float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return False
    return True


def pumsam_rate_value(row: PumsamRow) -> float:
    value = row.get("할증%")
    if value is None or value == "":
        return 1.0
    try:
        number = float(value)
    except (TypeError, ValueError):
        return 1.0
    if number > 5:
        return number / 100.0
    return number


_SIZE_RE = re.compile(
    r"(?P<num>\d+(?:\.\d+)?)\s*(?P<unit>mm2|cm2|kva|kv|mm|cm|m|a|p|w)\s*(?P<qual>이하|초과|미만|이상)?",
    re.I,
)
_PREFERRED_SIZE_UNITS = {"mm2", "cm2", "kva", "kv", "a", "p", "mm"}


_SPEC_PREFIXES = (
    "난연성",
    "난연",
    "분체도장",
    "아연도강전선관",
    "아연도강",
    "아연도",
    "후강전선관",
    "후강",
    "박강",
    "합성수지제전선관",
    "합성수지전선관",
    "합성수지제",
    "합성수지",
    "경질비닐전선관",
    "경질비닐",
    "hivp전선관",
    "hivp",
    "hfix전선",
    "hfix",
    "hiv전선",
    "hiv",
    "hi관",
    "cd관",
    "pe관",
    "ef전선관",
    "hi",
    "cd",
    "pe",
    "ef",
)


def _compact_stem(text: str) -> str:
    text = (
        str(text or "")
        .replace("×", "x")
        .replace("✕", "x")
        .replace("＊", "x")
        .replace("*", "x")
        .replace("Ｘ", "x")
        .replace("ｘ", "x")
    )
    return "".join(ch for ch in text if not ch.isspace()).lower()


def _clean_stem(stem: str) -> str:
    """아연도·G·HI 처럼 같은 관을 가리키는 앞말만 걷어, 숫자 규격으로 맞춘다."""
    text = _compact_stem(stem).replace("-", "").replace("_", "")
    changed = True
    while text and changed:
        changed = False
        for token in _SPEC_PREFIXES:
            if text == token or text.startswith(token):
                text = text[len(token) :]
                changed = True
                break
            if text.endswith(token):
                text = text[: -len(token)]
                changed = True
                break
    if text == "g":
        return ""
    return text


_WH_RE = re.compile(
    r"(?i)(?:w|ｗ)\s*(?P<w>\d+(?:\.\d+)?)\s*[x×*]\s*(?:h|ｈ)\s*(?P<h>\d+(?:\.\d+)?)"
)


def _wh_area(text: str) -> float | None:
    """W300 × H100 처럼 트레이 폭×높이이면 단면적(㎟)으로 본다."""
    match = _WH_RE.search(str(text or ""))
    if not match:
        return None
    return float(match.group("w")) * float(match.group("h"))


def spec_match_parts(spec: Any) -> tuple[float | None, str, str, str]:
    """규격에서 (숫자, 단위, 이하/초과, 나머지 말)을 뽑는다. 검색용 단위는 mm2."""
    text = lookup_measure(display_spec(spec or ""))
    area = _wh_area(text)
    matches = list(_SIZE_RE.finditer(text))
    if area is not None and not any((match.group("unit") or "").lower() == "mm2" for match in matches):
        return area, "mm2", "", ""
    if not matches:
        return None, "", "", _compact_stem(text)

    chosen = None
    for match in matches:
        if (match.group("qual") or "") == "이하":
            chosen = match
            break
    if chosen is None:
        for match in reversed(matches):
            if match.group("unit").lower() in _PREFERRED_SIZE_UNITS:
                chosen = match
                break
    if chosen is None:
        chosen = matches[-1]

    size = float(chosen.group("num"))
    unit = chosen.group("unit").lower()
    qual = (chosen.group("qual") or "").strip()
    stem = _compact_stem(text[: chosen.start()] + text[chosen.end() :])
    return size, unit, qual, stem


def _match_pumsam_ceiling(name: Any, spec: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    """딱 맞는 규격이 없으면, 같은 명칭의 가장 작은 「이하」 구간을 쓴다. 4㎟ → 6㎟ 이하."""
    name_key = lookup_key(name, "")
    item_size, item_unit, _, item_stem = spec_match_parts(spec)
    if not name_key or item_size is None or not item_unit:
        return []

    by_threshold: dict[float, dict[str, list[PumsamRow]]] = {}
    for row in rows:
        if name_key not in pumsam_name_keys(row):
            continue
        book_size, book_unit, book_qual, book_stem = spec_match_parts(row.get("규격"))
        if book_qual != "이하" or book_size is None or book_unit != item_unit:
            continue
        if _clean_stem(book_stem) != _clean_stem(item_stem):
            continue
        if book_size + 1e-9 < item_size:
            continue
        spec_key = lookup_key(row.get("명칭"), row.get("규격"))
        by_threshold.setdefault(book_size, {}).setdefault(spec_key, []).append(row)

    if not by_threshold:
        return []
    best_size = min(by_threshold)
    groups = by_threshold[best_size]
    return max(groups.values(), key=len)


def _match_pumsam_same_size(name: Any, spec: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    """같은 명칭에서 숫자·단위가 같고, 아연도/G/HI 앞말만 다른 규격을 같은 품으로 본다."""
    name_key = lookup_key(name, "")
    item_size, item_unit, _, item_stem = spec_match_parts(spec)
    if not name_key or item_size is None or not item_unit:
        return []
    item_core = _clean_stem(item_stem)
    item_stem_compact = _compact_stem(item_stem)

    by_spec: dict[str, list[PumsamRow]] = {}
    for row in rows:
        if name_key not in pumsam_name_keys(row):
            continue
        book_size, book_unit, book_qual, book_stem = spec_match_parts(row.get("규격"))
        if book_qual == "이하" or book_size is None or book_unit != item_unit:
            continue
        if abs(book_size - item_size) > 1e-9:
            continue
        if _clean_stem(book_stem) != item_core:
            continue
        spec_key = lookup_key(row.get("명칭"), row.get("규격"))
        by_spec.setdefault(spec_key, []).append(row)
    if not by_spec:
        return []

    def _rank(group: list[PumsamRow]) -> tuple[int, int]:
        _, _, _, stem = spec_match_parts(group[0].get("규격"))
        same_words = 1 if _compact_stem(stem) == item_stem_compact else 0
        return (same_words, len(group))

    return max(by_spec.values(), key=_rank)


def _match_pumsam_unique_item(name: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    """그 이름으로 규격이 한 가지뿐이면 그 인부들을 쓴다. 수평도체처럼 원표에 항목이 하나일 때."""
    name_key = lookup_key(name, "")
    if not name_key:
        return []
    by_spec: dict[str, list[PumsamRow]] = {}
    for row in rows:
        if name_key not in pumsam_name_keys(row):
            continue
        spec_key = lookup_key("", row.get("규격"))
        by_spec.setdefault(spec_key, []).append(row)
    if len(by_spec) != 1:
        return []
    return next(iter(by_spec.values()))


def _preferred_match_group(name: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    """짧은명칭과 공식 명칭이 같이 맞으면 검색한 이름 쪽만 남긴다."""
    kept = [row for row in rows if has_pumsam_qty(row)]
    if not kept:
        return []
    by_name: dict[str, list[PumsamRow]] = {}
    for row in kept:
        by_name.setdefault(lookup_key(row.get("명칭"), ""), []).append(row)
    if len(by_name) == 1:
        return kept
    query = lookup_key(name, "")
    if query in by_name:
        return by_name[query]
    return max(by_name.values(), key=lambda group: (len(group), -len(str(group[0].get("명칭") or ""))))


def strip_place_clone_rows(rows: list[PumsamRow]) -> list[PumsamRow]:
    """원표와 같은 품값인 노출·매입·지중·직매 복제 행은 품셈표에서 뺀다."""
    bases: dict[str, PumsamRow] = {}
    for row in rows:
        name = str(row.get("명칭") or "")
        base, suffix = split_place_name(name)
        if suffix:
            continue
        key = f"{lookup_key(name, row.get('규격'))}|{row.get('노무명칭')}"
        bases[key] = row
    kept: list[PumsamRow] = []
    for row in rows:
        name = str(row.get("명칭") or "")
        base, suffix = split_place_name(name)
        if suffix not in PLACE_SURCHARGE_RATES:
            kept.append(row)
            continue
        peer = bases.get(f"{lookup_key(base, row.get('규격'))}|{row.get('노무명칭')}")
        if peer is None:
            kept.append(row)
            continue
        if abs(pumsam_qty_value(peer) - pumsam_qty_value(row)) > 1e-9:
            kept.append(row)
            continue
        try:
            row_rate = int(float(row.get("할증%") or 100))
        except (TypeError, ValueError):
            kept.append(row)
            continue
        if row_rate != PLACE_SURCHARGE_RATES[suffix]:
            kept.append(row)
            continue
    return kept


def with_place_surcharge(rows: list[PumsamRow], suffix: str) -> list[PumsamRow]:
    rate = PLACE_SURCHARGE_RATES.get(suffix)
    if not rate:
        return rows
    copied: list[PumsamRow] = []
    for row in rows:
        item = dict(row)
        item["할증%"] = rate
        copied.append(item)
    return copied


def _match_pumsam_candidates(names: list[str], spec: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    for candidate in names:
        key = lookup_key(candidate, spec)
        exact = [
            row
            for row in rows
            if lookup_key(row.get("명칭"), row.get("규격")) == key
            or lookup_key(row.get("짧은명칭"), row.get("규격")) == key
        ]
        if exact:
            return _preferred_match_group(candidate, exact)
        same_size = _match_pumsam_same_size(candidate, spec, rows)
        if same_size:
            return _preferred_match_group(candidate, same_size)
        ceiling = _match_pumsam_ceiling(candidate, spec, rows)
        if ceiling:
            return _preferred_match_group(candidate, ceiling)
        unique = _match_pumsam_unique_item(candidate, rows)
        if unique:
            return [row for row in unique if has_pumsam_qty(row)]
    compact = lookup_key(names[0] if names else "", "").lower()
    if "tray" in compact or "트레이" in compact:
        for candidate in names:
            ceiling = _match_pumsam_ceiling(candidate, "1 mm2", rows)
            if ceiling:
                return _preferred_match_group(candidate, ceiling)
    return []


def match_pumsam(name: Any, spec: Any, rows: list[PumsamRow]) -> list[PumsamRow]:
    """같은 명칭·규격의 인부 행을 모두 반환한다.

    글자가 똑같으면 그걸 쓴다. 후강전선관 = 강제전선관처럼 같은 품 묶음이면 그 이름으로도 찾는다.
    아연도 16 mm 와 16 mm / G 16 mm 처럼 앞말만 다르면 같은 크기로 맞춘다.
    그래도 없으면 품목 규격 이상인 가장 작은 「이하」 구간을 씁니다.
    단가대비표의 _노출/_지중/_직매는 원표 품을 찾고 할증만 붙입니다.
    세대분전반_노출처럼 원표에 접미사가 따로 있는 품은 그 줄을 그대로 씁니다.
    """
    if not lookup_key(name, spec):
        return []
    found = _match_pumsam_candidates(alias_names(name), spec, rows)
    base, suffix = split_place_name(name)
    if suffix not in PLACE_SURCHARGE_RATES:
        return found
    if found and any(split_place_name(row.get("명칭"))[1] == suffix for row in found):
        return found
    if not found:
        found = _match_pumsam_candidates(alias_names(base), spec, rows)
    return with_place_surcharge(found, suffix)


def matched_surcharge_percent(name: Any, spec: Any, rows: list[PumsamRow] | None) -> int:
    """공량산출서 할증% 칸. 품셈 숫자는 원표 값이고, 할증만 따로 적는다."""
    if rows:
        found = match_pumsam(name, spec, rows)
        if found:
            raw = found[0].get("할증%")
            try:
                number = float(raw)
            except (TypeError, ValueError):
                number = 100.0
            return int(round(number if number > 5 else number * 100.0))
    _base, suffix = split_place_name(name)
    return int(PLACE_SURCHARGE_RATES.get(suffix, 100))


def surcharge_percent_text(row: PumsamRow) -> str:
    """할증% 칸을 사람이 읽는 퍼센트 글자로. 품셈 숫자와 섞지 않는다."""
    raw = row.get("할증%")
    if raw is None or raw == "":
        percent = 100.0
    else:
        try:
            number = float(raw)
        except (TypeError, ValueError):
            percent = 100.0
        else:
            percent = number if number > 5 else number * 100.0
    if float(percent).is_integer():
        return str(int(percent))
    return str(percent)


def pumsam_surcharge_note(row: PumsamRow) -> str:
    """결과 품셈표 비고. 품이 없으면 적지 않는다."""
    if not has_pumsam_qty(row):
        return ""
    qty = pumsam_qty_value(row)
    return f"품셈 {qty:.3f} · 할증 {surcharge_percent_text(row)}%"


def labor_kind_text(row: PumsamRow) -> str:  # noqa: ARG001
    """일위대가 인부 규격 칸. 할증은 적지 않는다."""
    return "일반공사 직종"


def labor_names_text(rows: list[PumsamRow]) -> str:
    seen: list[str] = []
    for row in rows:
        job = str(row.get("노무명칭") or "").strip()
        if job and job not in seen:
            seen.append(job)
    return ", ".join(seen)


def pumsam_db_path(directory: Path | None = None, discipline: str | None = None) -> Path:
    return user_database_dir(directory) / pumsam_filename(discipline)


def bundled_pumsam_path(discipline: str | None = None) -> Path:
    disc = normalize_discipline(discipline)
    named = bundled_data_dir() / pumsam_filename(disc)
    if named.exists():
        return named
    if disc == ELECTRIC:
        shared = bundled_data_dir() / PUMSAM_DB_FILENAME
        if shared.exists():
            return shared
    return named


def _load_rows_from_path(path: Path) -> list[PumsamRow]:
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        for candidate in workbook.worksheets:
            if "품셈" in str(candidate.title):
                sheet = candidate
                break
        grid = fill_merged_values(sheet)
    finally:
        workbook.close()
    table = trim_grid(grid)
    if not table:
        return []
    header_idx = find_header_row(table)
    return rows_from_grid(trim_grid(table[header_idx:]))


def ensure_pumsam_database(directory: Path | None = None, discipline: str | None = None) -> Path:
    disc = normalize_discipline(discipline)
    dest = pumsam_db_path(directory, disc)
    if dest.exists():
        if dest.stat().st_size < 50_000:
            bundled = bundled_pumsam_path(disc)
            if bundled.exists() and bundled.stat().st_size > dest.stat().st_size:
                dest.write_bytes(bundled.read_bytes())
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    bundled = bundled_pumsam_path(disc)
    if bundled.exists():
        dest.write_bytes(bundled.read_bytes())
        return dest
    if disc == ELECTRIC:
        shared = user_database_dir(directory) / PUMSAM_DB_FILENAME
        if shared.exists() and shared != dest:
            dest.write_bytes(shared.read_bytes())
            return dest
        legacy = get_result_directory() / LEGACY_PUMSAM_DB_FILENAME
        if directory is not None:
            legacy_alt = Path(directory) / LEGACY_PUMSAM_DB_FILENAME
            if legacy_alt.exists():
                dest.write_bytes(legacy_alt.read_bytes())
                return dest
        if legacy.exists():
            dest.write_bytes(legacy.read_bytes())
            return dest
    save_pumsam_database(default_pumsam_rows(disc), directory, discipline=disc)
    return dest


def ensure_all_pumsam_databases(directory: Path | None = None) -> None:
    """전기·통신 표준품셈 파일이 없으면 씨앗으로 만든다. 아주 작은 옛 전기 씨앗은 새 베이스로 바꾼다."""
    for disc in (ELECTRIC, TELECOM):
        ensure_pumsam_database(directory, disc)


def _hi_spec(mm: int) -> str:
    return f"HI {mm} mm"


def _conduit_seed_rows(labor: str, extra_labor: str | None = None) -> list[PumsamRow]:
    """경질비닐전선관 씨앗. 전기자재라 품셈근거는 전기 5-1을 유지한다."""
    buried = {
        16: 0.050,
        22: 0.060,
        28: 0.080,
        36: 0.100,
        42: 0.130,
        54: 0.190,
        70: 0.280,
        82: 0.370,
        92: 0.450,
        104: 0.460,
    }
    exposed = {
        16: 0.060,
        22: 0.072,
        28: 0.096,
        36: 0.120,
        42: 0.156,
        54: 0.228,
        70: 0.336,
        82: 0.444,
        92: 0.540,
        104: 0.552,
    }
    rows: list[PumsamRow] = []
    for mm, value in buried.items():
        name = "경질비닐전선관_지중"
        spec = _hi_spec(mm)
        rows.append(
            {
                "검색키": lookup_key(name, spec),
                "명칭": name,
                "규격": spec,
                "단위": "M",
                "노무명칭": labor,
                "품셈": value,
                "할증%": 100,
                "품셈근거": "전기5-1",
            }
        )
    for mm, value in exposed.items():
        name = "경질비닐전선관_노출"
        spec = _hi_spec(mm)
        rows.append(
            {
                "검색키": lookup_key(name, spec),
                "명칭": name,
                "규격": spec,
                "단위": "M",
                "노무명칭": labor,
                "품셈": value,
                "할증%": 120,
                "품셈근거": "전기5-1",
            }
        )
    if extra_labor:
        rows.append(
            {
                "검색키": lookup_key("경질비닐전선관_노출", _hi_spec(104)),
                "명칭": "경질비닐전선관_노출",
                "규격": _hi_spec(104),
                "단위": "M",
                "노무명칭": extra_labor,
                "품셈": 0.001,
                "할증%": 120,
                "품셈근거": "전기5-1",
            }
        )
    return rows


def default_pumsam_rows(discipline: str | None = None) -> list[PumsamRow]:
    """선택한 파트의 씨앗 품셈. 전기와 통신을 섞지 않는다."""
    disc = normalize_discipline(discipline)
    if disc == ELECTRIC:
        return electric_pumsam_rows()
    return telecom_pumsam_rows()


def _row_to_dict(values: list[Any]) -> PumsamRow | None:
    name = values[1] if len(values) > 1 else None
    spec = values[2] if len(values) > 2 else None
    if name is None and spec is None:
        return None
    return {
        "검색키": values[0] or lookup_key(name, spec),
        "명칭": name,
        "규격": spec,
        "단위": values[3] if len(values) > 3 else None,
        "노무명칭": values[4] if len(values) > 4 else None,
        "품셈": values[5] if len(values) > 5 else None,
        "할증%": values[6] if len(values) > 6 else None,
        "품셈근거": values[7] if len(values) > 7 else None,
    }


_HEADER_LIKE_NAMES = {"명칭", "품명", "품목", "검색키", "키워드", "품목명"}
_GROUP_TITLES = {"공량산출", "품목", "비고"}


def _combined_header(table: SheetRows) -> tuple[list[Any], int]:
    """1~2단 헤더를 한 줄로 합치고, 데이터 시작 인덱스를 반환한다."""
    if not table:
        return [], 0
    header_idx = find_header_row(table)
    span = header_row_span(table, header_idx)
    row1 = table[header_idx]
    row2 = table[header_idx + 1] if span == 2 and header_idx + 1 < len(table) else []
    width = max(len(row1), len(row2))
    combined: list[Any] = []
    for col in range(width):
        top = row1[col] if col < len(row1) else None
        bottom = row2[col] if col < len(row2) else None
        bottom_token = normalize_header(bottom)
        top_token = normalize_header(top)
        if bottom_token and bottom_token not in _GROUP_TITLES:
            combined.append(bottom)
        elif top_token and top_token not in _GROUP_TITLES:
            combined.append(top)
        else:
            combined.append(bottom if bottom is not None else top)
    return combined, header_idx + span


def rows_from_grid(table: SheetRows) -> list[PumsamRow]:
    if not table:
        return []
    header, data_start = _combined_header(table)
    spec_idx = find_column_index(header, "규격")
    unit_idx = find_column_index(header, "단위")
    name_idx = None
    labor_idx = None
    pumsam_idx = None
    rate_idx = None
    ref_idx = None
    seen_name = False
    for i, cell in enumerate(header):
        token = normalize_header(cell)
        if token in {"명칭", "품명"}:
            if not seen_name:
                name_idx = i
                seen_name = True
            else:
                labor_idx = i
        elif token in {"노무명칭", "직종", "공사인원"}:
            labor_idx = i
        elif token == "품셈":
            pumsam_idx = i
        elif token in {"할증%", "할증", "노무할증%", "할증률"}:
            rate_idx = i
        elif token in {"품셈근거", "근거"}:
            ref_idx = i
        elif token in {"검색키", "키워드"}:
            continue
    if name_idx is None:
        name_idx = find_column_index(header, "명칭")

    parsed: list[PumsamRow] = []
    prev_name: Any = None
    prev_spec: Any = None
    prev_unit: Any = None
    for source in table[data_start:]:
        def pick(index: int | None) -> Any:
            if index is None or index >= len(source):
                return None
            value = source[index]
            if isinstance(value, str) and not value.strip():
                return None
            return value

        name = pick(name_idx)
        spec = pick(spec_idx)
        labor = pick(labor_idx)
        pumsam_value = pick(pumsam_idx)
        if name is None and spec is None:
            if (labor or pumsam_value is not None) and prev_name is not None:
                name, spec = prev_name, prev_spec
                unit_value = prev_unit
            else:
                continue
        else:
            if name is None:
                name = prev_name
            if spec is None:
                spec = prev_spec
            unit_value = pick(unit_idx)
            if unit_value is None:
                unit_value = prev_unit
            prev_name, prev_spec, prev_unit = name, spec, unit_value
        if normalize_header(name) in _HEADER_LIKE_NAMES:
            continue
        if normalize_header(pumsam_value) == "품셈":
            continue
        parsed.append(
            {
                "검색키": lookup_key(name, spec),
                "명칭": name,
                "규격": spec,
                "단위": unit_value if name is not None else prev_unit,
                "노무명칭": labor,
                "품셈": pumsam_value,
                "할증%": pick(rate_idx),
                "품셈근거": pick(ref_idx),
            }
        )
    return parsed


def pumsam_identity(row: PumsamRow) -> str:
    labor = str(row.get("노무명칭") or "").strip()
    return f"{lookup_key(row.get('명칭'), row.get('규격'))}|{labor}"


def merge_pumsam_rows(*groups: list[PumsamRow]) -> list[PumsamRow]:
    merged: dict[str, PumsamRow] = {}
    for group in groups:
        for row in group:
            key = pumsam_identity(row)
            if not key or key == "|":
                continue
            row = dict(row)
            row["검색키"] = lookup_key(row.get("명칭"), row.get("규격"))
            merged[key] = row
    return strip_place_clone_rows(list(merged.values()))


def load_pumsam_database(directory: Path | None = None, discipline: str | None = None) -> list[PumsamRow]:
    disc = normalize_discipline(discipline)
    path = ensure_pumsam_database(directory, disc)
    parsed = _load_rows_from_path(path) if path.exists() else []
    seed = default_pumsam_rows(disc)
    if not parsed:
        return seed
    return merge_pumsam_rows(seed, parsed)


def _align(kind: str) -> Alignment:
    return Alignment(horizontal=kind, vertical="center", wrap_text=False)


def _style_pumsam_sheet(sheet) -> None:
    last_row = max(sheet.max_row, 1)
    last_col = len(PUMSAM_DISPLAY_HEADERS)
    for row_idx in range(1, last_row + 1):
        sheet.row_dimensions[row_idx].height = PUMSAM_ROW_HEIGHT
        is_header = row_idx == 1
        for col_idx in range(1, last_col + 1):
            cell = sheet.cell(row_idx, col_idx)
            cell.font = GULIM_HEADER if is_header else GULIM
            cell.border = THIN
            cell.fill = SKY_BLUE if is_header else WHITE
            if is_header:
                cell.alignment = _align("center")
            else:
                cell.alignment = _align(PUMSAM_DATA_ALIGNS[col_idx - 1])
            if col_idx == 6 and not is_header and isinstance(cell.value, (int, float)):
                cell.number_format = "0.000"
    for col_idx, width in enumerate(PUMSAM_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(last_col)}{last_row}"


def _style_plain_sheet(sheet, widths: list[int]) -> None:
    last_row = max(sheet.max_row, 1)
    last_col = max(sheet.max_column, 1)
    for row_idx in range(1, last_row + 1):
        sheet.row_dimensions[row_idx].height = PUMSAM_ROW_HEIGHT
        for col_idx in range(1, last_col + 1):
            cell = sheet.cell(row_idx, col_idx)
            cell.font = GULIM
            cell.alignment = _align("left")
    for col_idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width


def _pumsam_item_key(row: PumsamRow) -> tuple[str, str]:
    return lookup_key(row.get("명칭"), ""), lookup_key("", row.get("규격"))


def pumsam_row_sort_key(row: PumsamRow) -> tuple:
    return (
        pumsam_ref_sort_key(row.get("품셈근거")),
        str(row.get("명칭") or ""),
        str(row.get("규격") or ""),
        str(row.get("노무명칭") or ""),
    )


def _sort_pumsam_rows(rows: list[PumsamRow]) -> list[PumsamRow]:
    return sorted(rows, key=pumsam_row_sort_key)


def fill_pumsam_search_sheet(sheet, rows: list[PumsamRow]) -> None:
    """검색·원표 대조용 품셈표. 키워드·명칭·규격·단위·노무명칭·품셈·할증%·품셈근거."""
    sheet.title = PUMSAM_SHEET_NAME
    sheet.append(list(PUMSAM_DISPLAY_HEADERS))
    for row, hide_item in grouped_pumsam_rows(rows):
        sheet.append(_sheet_values(row, hide_item=hide_item))
    _style_pumsam_sheet(sheet)


def grouped_pumsam_rows(rows: list[PumsamRow]) -> list[tuple[PumsamRow, bool]]:
    """같은 명칭·규격의 다음 인부 행은 키워드·명칭만 비운다. 규격은 그대로."""
    grouped: list[tuple[PumsamRow, bool]] = []
    prev_key: tuple[str, str] | None = None
    for row in _sort_pumsam_rows(rows):
        key = _pumsam_item_key(row)
        hide_item = prev_key == key and bool(key[0])
        grouped.append((row, hide_item))
        prev_key = key
    return grouped


def _sheet_values(row: PumsamRow, *, hide_item: bool) -> list[Any]:
    spec = display_spec(row.get("규격") or "") or row.get("규격")
    return [
        None if hide_item else display_keyword(row.get("명칭"), row.get("규격")),
        None if hide_item else row.get("명칭"),
        spec,
        row.get("단위"),
        row.get("노무명칭"),
        row.get("품셈"),
        row.get("할증%"),
        row.get("품셈근거"),
    ]


def save_pumsam_database(
    rows: list[PumsamRow],
    directory: Path | None = None,
    discipline: str | None = None,
) -> Path:
    disc = normalize_discipline(discipline)
    rows = strip_place_clone_rows(list(rows))
    path = pumsam_db_path(directory, disc)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    fill_pumsam_search_sheet(sheet, rows)
    if disc == ELECTRIC:
        rules = workbook.create_sheet("적용기준")
        rules.append(["적용 기준"])
        for line in ELECTRIC_RULES:
            rules.append([line])
        trades = workbook.create_sheet("공종별인부")
        for guide_row in ELECTRIC_TRADE_GUIDE:
            trades.append(list(guide_row))
        _style_plain_sheet(rules, [90])
        _style_plain_sheet(trades, [28, 14, 28, 12, 18, 12, 48])
        trades.row_dimensions[1].height = PUMSAM_ROW_HEIGHT
        for col_idx in range(1, 8):
            cell = trades.cell(1, col_idx)
            cell.font = GULIM_HEADER
            cell.fill = SKY_BLUE
            cell.alignment = _align("center")
            cell.border = THIN
    else:
        rules = workbook.create_sheet("적용기준")
        rules.append(["적용 기준"])
        for line in TELECOM_RULES:
            rules.append([line])
        trades = workbook.create_sheet("공종별인부")
        for guide_row in TELECOM_TRADE_GUIDE:
            trades.append(list(guide_row))
        _style_plain_sheet(rules, [90])
        _style_plain_sheet(trades, [28, 14, 28, 12, 18, 12, 48])
        for col_idx in range(1, 8):
            cell = trades.cell(1, col_idx)
            cell.font = GULIM_HEADER
            cell.fill = SKY_BLUE
            cell.alignment = _align("center")
            cell.border = THIN
    workbook.save(path)
    workbook.close()
    return path


def import_pumsam_file(source_path: Path) -> list[PumsamRow]:
    path = Path(source_path)
    if not is_allowed_excel(path):
        raise ValueError("xlsx 또는 xlsm 파일만 읽을 수 있습니다.")
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        for candidate in workbook.worksheets:
            if "품셈" in str(candidate.title):
                sheet = candidate
                break
        grid = fill_merged_values(sheet)
    finally:
        workbook.close()
    table = trim_grid(grid)
    if not table:
        return []
    header_idx = find_header_row(table)
    return rows_from_grid(trim_grid(table[header_idx:]))
