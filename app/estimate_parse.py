"""내역서 엑셀 읽기. 원본 파일은 저장하지 않는다."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from app.merge_parse import SheetRows, fill_merged_values, trim_grid
from app.paths import is_allowed_excel

HEADER_ALIASES = {
    "명칭": ("명칭", "품명", "품목", "자재명", "항목"),
    "규격": ("규격", "사양", "규격/사양"),
    "단위": ("단위", "단위명"),
    "수량": ("수량", "설계수량", "물량", "결정수량", "계약수량", "설계물량"),
}

HEADER_HINTS = ("명칭", "품명", "규격", "단위", "수량", "단가")
SECTION_NAME = re.compile(r"^\s*\d+\s*[\.．]")


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace(" ", "").replace("\n", "").strip()


def find_column_index(header_row: list[Any], field: str) -> int | None:
    aliases = HEADER_ALIASES.get(field, (field,))
    normalized_aliases = {normalize_header(alias) for alias in aliases}
    for idx, cell in enumerate(header_row):
        if normalize_header(cell) in normalized_aliases:
            return idx
    return None


def find_quantity_column(header_row: list[Any]) -> int | None:
    idx = find_column_index(header_row, "수량")
    if idx is not None:
        return idx
    for i, cell in enumerate(header_row):
        name = normalize_header(cell)
        if "수량" in name and "단가" not in name and "금액" not in name:
            return i
    return None


def find_header_row(rows: SheetRows) -> int:
    for i, row in enumerate(rows):
        normalized = [normalize_header(c) for c in row]
        hits = sum(1 for hint in HEADER_HINTS if any(hint in cell for cell in normalized))
        if hits >= 2:
            return i
    return 0


def is_section_row(name: Any, spec: Any, unit: Any) -> bool:
    """'1. 옥외전기공사' 같은 공종 제목 행."""
    name_text = str(name).strip() if name is not None else ""
    if not name_text:
        return True
    if spec is None and unit is None and SECTION_NAME.match(name_text):
        return True
    if spec is None and unit is None and name_text.endswith("공사") and "_" not in name_text:
        return True
    return False


def lookup_key(name: Any, spec: Any) -> str:
    """품셈표 검색용. 공백을 없애 명칭+규격을 붙인다."""
    left = "" if name is None else str(name)
    right = "" if spec is None else str(spec)
    return "".join(ch for ch in f"{left}{right}" if not ch.isspace())


def read_workbook_first_sheet(source_path: Path) -> SheetRows:
    path = Path(source_path)
    if not path.exists():
        raise FileNotFoundError(f"원본 파일을 찾을 수 없습니다: {path}")
    if not is_allowed_excel(path):
        raise ValueError("xlsx 또는 xlsm 파일만 읽을 수 있습니다.")

    from openpyxl import load_workbook

    with path.open("rb") as handle:
        workbook = load_workbook(filename=handle, data_only=True, keep_vba=False)
        try:
            if not workbook.worksheets:
                raise ValueError("엑셀에 시트가 없습니다.")
            sheet = workbook.worksheets[0]
            for candidate in workbook.worksheets:
                title = str(candidate.title)
                if "내역" in title or "단가" in title:
                    sheet = candidate
                    break
            grid = fill_merged_values(sheet)
        finally:
            workbook.close()

    cleaned = trim_grid(grid)
    if not cleaned:
        raise ValueError("내역서에 읽을 수 있는 데이터가 없습니다.")
    header_idx = find_header_row(cleaned)
    table = trim_grid(cleaned[header_idx:])
    if not table:
        raise ValueError("내역서 헤더(명칭/규격/단위/수량)를 찾지 못했습니다.")
    return table


def read_named_sheet_rows(source_path: Path, name_contains: str) -> SheetRows | None:
    """통합 엑셀에 품셈표 시트가 있으면 그 표를 반환한다."""
    from openpyxl import load_workbook

    path = Path(source_path)
    with path.open("rb") as handle:
        workbook = load_workbook(filename=handle, data_only=True, keep_vba=False)
        try:
            sheet = None
            for candidate in workbook.worksheets:
                if name_contains in str(candidate.title):
                    sheet = candidate
                    break
            if sheet is None:
                return None
            grid = fill_merged_values(sheet)
        finally:
            workbook.close()
    cleaned = trim_grid(grid)
    if not cleaned:
        return None
    header_idx = find_header_row(cleaned)
    return trim_grid(cleaned[header_idx:])


# 이전 이름 호환
read_unit_price_table = read_workbook_first_sheet
