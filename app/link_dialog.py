"""미매칭 품목에 품셈 항목을 직접 연결하는 창."""

from __future__ import annotations

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QDialog,
    QDialogButtonBox,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QVBoxLayout,
)


def distinct_book_items(pumsam_rows: list) -> list[tuple[str, str, str, str]]:
    """(명칭, 규격, 인부 요약, 근거) — 품셈 수치가 있는 항목만, 중복 없이."""
    found: dict[tuple[str, str], list[str]] = {}
    refs: dict[tuple[str, str], str] = {}
    for row in pumsam_rows:
        name = str(row.get("명칭") or "").strip()
        spec = str(row.get("규격") or "").strip()
        qty = row.get("품셈")
        if not name or not isinstance(qty, (int, float)) or qty <= 0:
            continue
        key = (name, spec)
        job = str(row.get("노무명칭") or "")
        found.setdefault(key, [])
        if job and job not in found[key]:
            found[key].append(job)
        refs.setdefault(key, str(row.get("품셈근거") or ""))
    return [(n, s, ", ".join(j), refs[(n, s)]) for (n, s), j in sorted(found.items())]


class LinkDialog(QDialog):
    def __init__(self, parent, item_label: str, pumsam_rows: list) -> None:
        super().__init__(parent)
        self.setWindowTitle("품셈 연결")
        self.resize(640, 480)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(f"'{item_label}' 에 적용할 품셈 항목을 고르세요. (다음부터 자동 적용됩니다)"))
        self.search = QLineEdit()
        self.search.setPlaceholderText("품셈 명칭·규격 검색 (띄어쓰기로 여러 낱말)")
        self.search.textChanged.connect(self._refill)
        layout.addWidget(self.search)
        self.list = QListWidget()
        layout.addWidget(self.list, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)
        self._items = distinct_book_items(pumsam_rows)
        self._refill()

    def _refill(self) -> None:
        tokens = [t for t in self.search.text().lower().split() if t]
        self.list.clear()
        shown = 0
        for name, spec, jobs, ref in self._items:
            text = f"{name} | {spec} | {jobs} | {ref}"
            if all(t in text.lower() for t in tokens):
                entry = QListWidgetItem(text)
                entry.setData(Qt.ItemDataRole.UserRole, (name, spec))
                self.list.addItem(entry)
                shown += 1
                if shown >= 400:
                    break

    def chosen(self) -> tuple[str, str] | None:
        current = self.list.currentItem()
        return current.data(Qt.ItemDataRole.UserRole) if current is not None else None


class ColumnMapDialog(QDialog):
    """낯선 단가대비표의 어느 열이 명칭·규격·단위·단가인지 직접 고르는 창."""

    def __init__(self, parent, preview: list[list]) -> None:
        super().__init__(parent)
        from PyQt6.QtWidgets import QComboBox, QFormLayout, QSpinBox, QTableWidget, QTableWidgetItem
        from openpyxl.utils import get_column_letter

        self.setWindowTitle("단가대비표 열 지정")
        self.resize(760, 520)
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel("이 양식은 자동으로 읽지 못했습니다. 아래 미리보기를 보고 어느 열이 무엇인지 골라 주세요."))
        width = max((len(r) for r in preview), default=1)
        table = QTableWidget(len(preview), width)
        table.setHorizontalHeaderLabels([get_column_letter(i + 1) for i in range(width)])
        for r, row in enumerate(preview):
            for c in range(width):
                value = row[c] if c < len(row) else None
                table.setItem(r, c, QTableWidgetItem("" if value is None else str(value)))
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(table, 1)
        form = QFormLayout()
        self.first_row = QSpinBox()
        self.first_row.setRange(1, 999)
        self.first_row.setValue(min(3, max(1, len(preview))))
        form.addRow("자료가 시작하는 행 번호", self.first_row)
        letters = [get_column_letter(i + 1) for i in range(max(width, 4))]

        def combo(default: int, optional: bool = False) -> QComboBox:
            box = QComboBox()
            if optional:
                box.addItem("(없음)", None)
            for i, letter in enumerate(letters):
                box.addItem(f"{letter} 열", i + 1)
            box.setCurrentIndex(min(default + (1 if optional else 0), box.count() - 1))
            return box

        self.name_box = combo(0)
        self.spec_box = combo(1, optional=True)
        self.unit_box = combo(2, optional=True)
        self.price_box = combo(3)
        form.addRow("명칭(품명)", self.name_box)
        form.addRow("규격", self.spec_box)
        form.addRow("단위", self.unit_box)
        form.addRow("단가(재료비)", self.price_box)
        layout.addLayout(form)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def mapping(self) -> dict:
        return {
            "first_data_row": self.first_row.value(),
            "name_col": self.name_box.currentData(),
            "spec_col": self.spec_box.currentData(),
            "unit_col": self.unit_box.currentData(),
            "price_col": self.price_box.currentData(),
        }
