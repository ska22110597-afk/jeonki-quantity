"""단가대비표 → 일위대가 → 내역서 → 공량산출서 파이프라인."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.estimate_parse import EstimateSheet, load_estimate_sheet
from app.excel_io import (
    BODY_FONT,
    CENTER,
    COMPARE_SHEET_NAME,
    ESTIMATE_SHEET_NAME,
    HEADER_FONT,
    ILWIDAE_SHEET_NAME,
    LEFT,
    MONEY_FORMAT,
    PUMSAM_DATA_START,
    QUANTITY_SHEET_NAME,
    RIGHT,
    SECTION_FONT,
    _apply_sheet_look,
    _set_cell,
    _write_estimate_sheet,
    _write_pumsam_sheet,
    _write_quantity_sheet,
    collect_pumsam_rows,
)
from app.ilwidae import IlwidaeBlock, write_ilwidae_sheet
from app.items import LineItem, parse_line_items
from app.paths import assert_safe_save, build_result_path
from app.pumsam import PUMSAM_SHEET_NAME, PumsamRow
from app.wages import WAGES_SHEET_NAME, WageRow, load_wages, save_wages


def _write_compare_sheet(sheet: Worksheet, estimate: EstimateSheet) -> None:
    sheet.title = COMPARE_SHEET_NAME
    _write_estimate_sheet(sheet, estimate)
    if sheet["A1"].value in (None, "[내역서 ]"):
        sheet["A1"].value = "단 가 대 비 표"


def _write_wages_sheet(sheet: Worksheet, rows: list[WageRow]) -> None:
    sheet.title = WAGES_SHEET_NAME
    _set_cell(sheet, 1, 1, "직종", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, 1, 2, "노임단가", font=HEADER_FONT, align=CENTER)
    _set_cell(sheet, 1, 3, "비고", font=HEADER_FONT, align=CENTER)
    last = 1
    for offset, row in enumerate(rows):
        excel_row = offset + 2
        last = excel_row
        _set_cell(sheet, excel_row, 1, row.get("직종"), font=BODY_FONT)
        _set_cell(
            sheet,
            excel_row,
            2,
            row.get("노임단가"),
            font=BODY_FONT,
            align=RIGHT,
            number_format=MONEY_FORMAT,
        )
        _set_cell(sheet, excel_row, 3, row.get("비고"), font=BODY_FONT)
    _apply_sheet_look(sheet, last, 3)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 40


def _write_generated_estimate(
    sheet: Worksheet,
    items: list[LineItem],
    blocks: list[IlwidaeBlock],
) -> EstimateSheet:
    """일위대가 단가 × 단가대비표 수량으로 내역서를 만든다."""
    sheet.title = ESTIMATE_SHEET_NAME
    _set_cell(sheet, 1, 1, "[내역서 ]", font=BODY_FONT)
    for col in range(2, 14):
        _set_cell(sheet, 1, col, None, border=False)
    headers = {
        1: "명칭",
        2: "규격",
        3: "단위",
        4: "수량",
        5: "재료비",
        7: "노무비",
        9: "경비",
        11: "합계",
        13: "비고",
    }
    for col in range(1, 14):
        _set_cell(sheet, 2, col, headers.get(col), font=HEADER_FONT, align=CENTER)
        _set_cell(sheet, 3, col, None, font=HEADER_FONT, align=CENTER)
    for col, text in ((5, "단가"), (6, "금액"), (7, "단가"), (8, "금액"), (9, "단가"), (10, "금액"), (11, "단가"), (12, "금액")):
        _set_cell(sheet, 3, col, text, font=HEADER_FONT, align=CENTER)
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet.merge_cells("C2:C3")
    sheet.merge_cells("D2:D3")
    sheet.merge_cells("E2:F2")
    sheet.merge_cells("G2:H2")
    sheet.merge_cells("I2:J2")
    sheet.merge_cells("K2:L2")
    sheet.merge_cells("M2:M3")

    block_by_key = {block.item.key: block for block in blocks}
    filled: list[list[Any]] = [[None] * 13 for _ in range(3)]
    filled[0][0] = "[내역서 ]"
    filled[1][0] = "명칭"
    filled[1][1] = "규격"
    filled[1][2] = "단위"
    filled[1][3] = "수량"
    filled[1][4] = "재료비"
    filled[2][4] = "단가"
    filled[2][5] = "금액"

    excel_row = 4
    work_items = [item for item in items if not item.section]
    if work_items:
        _set_cell(sheet, excel_row, 1, "1. 전기공사", font=SECTION_FONT)
        for col in range(2, 14):
            _set_cell(sheet, excel_row, col, None)
        filled.append(["1. 전기공사"] + [None] * 12)
        excel_row += 1

    for item in items:
        if item.section:
            continue
        block = block_by_key.get(item.key)
        _set_cell(sheet, excel_row, 1, item.name, font=BODY_FONT)
        _set_cell(sheet, excel_row, 2, item.spec, font=BODY_FONT)
        _set_cell(sheet, excel_row, 3, item.unit, font=BODY_FONT, align=CENTER)
        _set_cell(
            sheet,
            excel_row,
            4,
            f"='{COMPARE_SHEET_NAME}'!D{item.excel_row}",
            font=BODY_FONT,
            align=RIGHT,
            number_format=MONEY_FORMAT,
        )
        _set_cell(
            sheet,
            excel_row,
            5,
            f"='{COMPARE_SHEET_NAME}'!E{item.excel_row}",
            font=BODY_FONT,
            align=RIGHT,
            number_format=MONEY_FORMAT,
        )
        _set_cell(sheet, excel_row, 6, f"=D{excel_row}*E{excel_row}", font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        labor_ref = f"='{ILWIDAE_SHEET_NAME}'!H{block.sum_row}" if block else 0
        _set_cell(sheet, excel_row, 7, labor_ref, font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 8, f"=D{excel_row}*G{excel_row}", font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 9, 0, font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 10, f"=D{excel_row}*I{excel_row}", font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 11, f"=E{excel_row}+G{excel_row}+I{excel_row}", font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 12, f"=F{excel_row}+H{excel_row}+J{excel_row}", font=BODY_FONT, align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, excel_row, 13, None)
        filled.append(
            [item.name, item.spec, item.unit, item.qty, item.material_price, None, None, None, None, None, None, None, None]
        )
        excel_row += 1

    _apply_sheet_look(sheet, max(excel_row - 1, 3), 13)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 8
    sheet.column_dimensions["D"].width = 10
    for col in range(5, 14):
        from openpyxl.utils import get_column_letter

        sheet.column_dimensions[get_column_letter(col)].width = 12
    return EstimateSheet(filled=filled, raw=filled, merges=[])


def _load_optional(path: Path | None) -> EstimateSheet | None:
    if path is None:
        return None
    return load_estimate_sheet(path)


def build_result_workbook(
    *,
    compare: EstimateSheet | None,
    estimate: EstimateSheet | None,
    pumsam_rows: list[PumsamRow],
    wage_rows: list[WageRow],
) -> Workbook:
    workbook = Workbook()
    first = workbook.active
    items: list[LineItem] = []
    blocks: list[IlwidaeBlock] = []

    if compare is not None:
        _write_compare_sheet(first, compare)
        items = parse_line_items(compare)
    elif estimate is not None:
        first.title = ESTIMATE_SHEET_NAME
        _write_estimate_sheet(first, estimate)
    else:
        first.title = ESTIMATE_SHEET_NAME

    if compare is not None and items:
        ilwidae = workbook.create_sheet(ILWIDAE_SHEET_NAME)
        blocks = write_ilwidae_sheet(ilwidae, items, pumsam_rows, wage_rows, compare_sheet=COMPARE_SHEET_NAME)
        generated = workbook.create_sheet(ESTIMATE_SHEET_NAME)
        estimate = _write_generated_estimate(generated, items, blocks)
    elif estimate is not None and first.title != ESTIMATE_SHEET_NAME:
        copied = workbook.create_sheet(ESTIMATE_SHEET_NAME)
        _write_estimate_sheet(copied, estimate)

    pumsam_last = PUMSAM_DATA_START + len(pumsam_rows) - 1 if pumsam_rows else PUMSAM_DATA_START
    if estimate is not None:
        qty = workbook.create_sheet(QUANTITY_SHEET_NAME)
        _write_quantity_sheet(qty, estimate, pumsam_last)

    pumsam_sheet = workbook.create_sheet(PUMSAM_SHEET_NAME)
    _write_pumsam_sheet(pumsam_sheet, pumsam_rows)
    wages_sheet = workbook.create_sheet(WAGES_SHEET_NAME)
    _write_wages_sheet(wages_sheet, wage_rows)

    if first.title == ESTIMATE_SHEET_NAME and compare is not None:
        # 단가대비표가 첫 시트. 내역서만 있던 경우 첫 시트는 내역서.
        pass
    return workbook


def run_pipeline(
    *,
    dest_dir: Path | None = None,
    unit_price_path: Path | None = None,
    ilwidae_path: Path | None = None,
    estimate_path: Path | None = None,
    extra_pumsam_path: Path | None = None,
    db_dir: Path | None = None,
    estimate_rows: list[list[Any]] | None = None,
) -> Path:
    database_dir = db_dir if db_dir is not None else dest_dir
    primary = unit_price_path or ilwidae_path or estimate_path
    if primary is None and estimate_rows is None:
        raise ValueError("단가대비표, 일위대가, 내역서 중 하나를 놓아 주세요.")

    compare = _load_optional(Path(unit_price_path) if unit_price_path else None)
    dropped_ilwidae = _load_optional(Path(ilwidae_path) if ilwidae_path else None)
    if estimate_rows is not None:
        estimate = EstimateSheet(filled=estimate_rows, raw=estimate_rows, merges=[])
    else:
        estimate = _load_optional(Path(estimate_path) if estimate_path else None)

    # 일위대가 파일만 있으면 그 품목으로 단가대비표를 대신한다.
    if compare is None and dropped_ilwidae is not None and estimate is None:
        compare = dropped_ilwidae

    source_for_pumsam = unit_price_path or estimate_path or ilwidae_path
    pumsam_rows = collect_pumsam_rows(
        Path(source_for_pumsam) if source_for_pumsam else None,
        database_dir,
        extra_pumsam_path,
        db_dir=database_dir,
    )
    wage_rows = load_wages(database_dir)
    save_wages(wage_rows, database_dir)

    dest = build_result_path(dest_dir)
    if primary is not None:
        assert_safe_save(Path(primary), dest)

    workbook = build_result_workbook(
        compare=compare,
        estimate=estimate,
        pumsam_rows=pumsam_rows,
        wage_rows=wage_rows,
    )
    try:
        workbook.save(dest)
    finally:
        workbook.close()
    return dest
