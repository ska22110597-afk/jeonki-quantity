"""체크박스로 고른 일위대가목록 품목으로 공정별 내역서를 만든다.

내역서 양식
- 맨 윗줄: 공정 제목(사용자가 입력)
- 품목 행: 재료비 단가 = 일위대가 자재 행 금액(자재 수량 할증 포함), 노무비 단가 = 일위대가 인부 금액 합
- 아래 집계: [배관 부속재] · [소모 잡자재] · 직종별 노무비 · [공구손료] · 합계
  합계 재료비 = 품목 재료비 + 부속재 + 잡자재, 합계 노무비 = 직종별 노무비 + 공구손료
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from app.discipline import normalize_discipline
from app.excel_io import (
    BODY_FONT,
    CENTER,
    COMPARE_SHEET_NAME,
    ESTIMATE_SHEET_NAME,
    FORM_ROW_HEIGHT,
    HEADER_FONT,
    ILWIDAE_LIST_SHEET_NAME,
    ILWIDAE_SHEET_NAME,
    LEFT,
    RIGHT,
    _apply_sheet_look,
    _set_cell,
    collect_pumsam_rows,
    write_title_banner,
)
from app.fitting_rules import load_fitting_excludes
from app.ilwidae import CD_FITTING_RATE, SUNDRY_RATE, TOOL_RATE, IlwidaeBlock
from app.paths import assert_safe_save, build_result_path
from app.wages import WAGES_SHEET_NAME, load_wages, save_wages

MONEY_FORMAT = "#,##0;-#,##0;;@"
QTY_FORMAT = "#,##0.##"
LAST_COL = 12
HELPER_START_COL = 14  # N 열부터는 계산용 보조 열(숨김)
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def trunc1(value: float) -> float:
    return math.floor(value * 10 + 1e-9) / 10


@dataclass
class IlwidaeEntry:
    """일위대가목록 한 줄(= 일위대가 호표 하나)."""

    index: int
    name: str
    spec: str
    unit: str
    ho_no: int
    material_row: int
    labor_rows: dict[str, list[int]]  # 직종 → 일위대가 시트의 인부 행들
    material_unit: float  # 자재 수량(할증 포함) × 단가
    labor_unit_by_job: dict[str, float]
    fitting_rate: float
    has_sundry: bool
    pumsam_ref: str = ""

    @property
    def labor_unit(self) -> float:
        return sum(self.labor_unit_by_job.values())

    @property
    def label(self) -> str:
        return f"{self.name} {self.spec}".strip()


@dataclass
class EstimateLine:
    entry: IlwidaeEntry
    qty: float
    remark: str = ""


@dataclass
class EstimateSection:
    title: str
    lines: list[EstimateLine] = field(default_factory=list)


@dataclass
class EstimateSession:
    workbook: Workbook
    entries: list[IlwidaeEntry]
    discipline: str
    source_path: Path | None = None
    estimate_sheets: list[str] = field(default_factory=list)  # 앞서 저장할 때 붙인 내역서 시트


def extract_entries(workbook: Workbook, blocks: list[IlwidaeBlock], wage_rows: list[dict]) -> list[IlwidaeEntry]:
    wages = {str(r.get("직종")): float(r.get("노임단가") or 0) for r in wage_rows}
    compare = workbook[COMPARE_SHEET_NAME]
    entries: list[IlwidaeEntry] = []
    for block in blocks:
        item = block.item
        col = (item.price_col if item.price_col is not None else 11) + 1
        price = compare.cell(item.excel_row, col).value
        price = float(price) if isinstance(price, (int, float)) else 0.0
        labor_rows: dict[str, list[int]] = {}
        labor_unit: dict[str, float] = {}
        for job, qty, row in zip(block.labor_jobs, block.labor_qtys, block.labor_rows):
            labor_rows.setdefault(job, []).append(row)
            labor_unit[job] = labor_unit.get(job, 0.0) + trunc1(wages.get(job, 0.0) * qty)
        entries.append(
            IlwidaeEntry(
                index=len(entries),
                name=str(item.name or ""),
                spec=str(item.spec or ""),
                unit=str(item.unit or ""),
                ho_no=block.ho_no,
                material_row=block.material_row,
                labor_rows=labor_rows,
                material_unit=trunc1(block.material_qty * price),
                labor_unit_by_job=labor_unit,
                fitting_rate=block.fitting_rate,
                has_sundry=block.has_sundry,
                pumsam_ref=block.pumsam_ref,
            )
        )
    return entries


def generate_session(
    unit_price_path: Path,
    discipline: str | None = None,
    db_dir: Path | None = None,
) -> EstimateSession:
    """단가대비표 → 일위대가·일위대가목록을 메모리에 만들고, 고를 수 있는 품목 목록을 돌려준다."""
    from app.pipeline import _load_optional, build_result_workbook

    disc = normalize_discipline(discipline)
    source = Path(unit_price_path)
    compare = _load_optional(source)
    pumsam_rows = collect_pumsam_rows(source, db_dir, None, db_dir=db_dir, discipline=disc)
    wage_rows = load_wages(db_dir)
    save_wages(wage_rows, db_dir)
    blocks: list[IlwidaeBlock] = []
    workbook = build_result_workbook(
        compare=compare,
        estimate=None,
        pumsam_rows=pumsam_rows,
        wage_rows=wage_rows,
        discipline=disc,
        mode="forward",
        fitting_excludes=load_fitting_excludes(db_dir),
        blocks_out=blocks,
    )
    entries = extract_entries(workbook, blocks, wage_rows)
    return EstimateSession(workbook=workbook, entries=entries, discipline=disc, source_path=source)


def _sheet_name(base: str, taken: set[str]) -> str:
    name = _BAD_SHEET_CHARS.sub("", base)[:31] or ESTIMATE_SHEET_NAME
    candidate = name
    n = 2
    while candidate in taken:
        suffix = f"({n})"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    return candidate


def _ref(row: int, col: str) -> str:
    return f"'{ILWIDAE_SHEET_NAME}'!{col}{row}"


def write_estimate_sheet(workbook: Workbook, name: str, section: EstimateSection) -> None:
    sheet = workbook.create_sheet(name)
    write_title_banner(sheet, "내   역   서", LAST_COL)
    heads = {1: "명  칭", 2: "규  격", 3: "단위", 4: "수량", 5: "재  료  비", 7: "노  무  비", 9: "경    비", 11: "계", 12: "비  고"}
    for col in range(1, LAST_COL + 1):
        _set_cell(sheet, 3, col, heads.get(col), font=HEADER_FONT, align=CENTER)
        _set_cell(sheet, 4, col, None, font=HEADER_FONT, align=CENTER)
    for col, text in ((5, "단  가"), (6, "금  액"), (7, "단  가"), (8, "금  액"), (9, "단  가"), (10, "금  액")):
        _set_cell(sheet, 4, col, text, font=HEADER_FONT, align=CENTER)
    for rng in ("A3:A4", "B3:B4", "C3:C4", "D3:D4", "E3:F3", "G3:H3", "I3:J3", "K3:K4", "L3:L4"):
        sheet.merge_cells(rng)

    lines = section.lines
    rates = sorted({ln.entry.fitting_rate for ln in lines if ln.entry.fitting_rate > 0})
    jobs: list[str] = []
    for ln in lines:
        for job, amount in ln.entry.labor_unit_by_job.items():
            if amount > 0 and job not in jobs:
                jobs.append(job)
    col_of: dict[Any, int] = {}
    next_col = HELPER_START_COL
    for rate in rates:
        col_of[("rate", rate)] = next_col
        next_col += 1
    col_of["sundry"] = next_col
    next_col += 1
    for job in jobs:
        col_of[("job", job)] = next_col
        next_col += 1
    helper_last = next_col - 1

    def letter(key: Any) -> str:
        return get_column_letter(col_of[key])

    title_row = 5
    _set_cell(sheet, title_row, 1, section.title, font=HEADER_FONT, align=LEFT)
    for col in range(2, LAST_COL + 1):
        _set_cell(sheet, title_row, col, None)
    sheet.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=LAST_COL)

    first = title_row + 1
    row = first
    for ln in lines:
        e = ln.entry
        _set_cell(sheet, row, 1, e.name)
        _set_cell(sheet, row, 2, e.spec)
        _set_cell(sheet, row, 3, e.unit, align=CENTER)
        _set_cell(sheet, row, 4, ln.qty, align=RIGHT, number_format=QTY_FORMAT)
        _set_cell(sheet, row, 5, f"={_ref(e.material_row, 'F')}", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, row, 6, f"=D{row}*E{row}", align=RIGHT, number_format=MONEY_FORMAT)
        job_cols = [letter(("job", j)) for j in jobs if e.labor_unit_by_job.get(j, 0) > 0]
        if job_cols:
            _set_cell(sheet, row, 7, "=" + "+".join(f"{c}{row}" for c in job_cols), align=RIGHT, number_format=MONEY_FORMAT)
            _set_cell(sheet, row, 8, f"=D{row}*G{row}", align=RIGHT, number_format=MONEY_FORMAT)
        else:
            _set_cell(sheet, row, 7, None)
            _set_cell(sheet, row, 8, None)
        _set_cell(sheet, row, 9, None)
        _set_cell(sheet, row, 10, None)
        _set_cell(sheet, row, 11, f"=F{row}+H{row}+J{row}", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, row, 12, ln.remark or None)
        for key, col in col_of.items():
            if key == "sundry":
                value = 1 if e.has_sundry else None
            elif key[0] == "rate":
                value = 1 if abs(e.fitting_rate - key[1]) < 1e-9 else None
            else:
                refs = e.labor_rows.get(key[1], [])
                value = "=" + "+".join(_ref(r, "H") for r in refs) if refs and e.labor_unit_by_job.get(key[1], 0) > 0 else None
            sheet.cell(row, col, value)
        row += 1
    last = max(row - 1, first)

    material_rows: list[int] = []
    labor_first = labor_last = None
    summary = row

    def blank(r: int) -> None:
        for c in range(1, LAST_COL + 1):
            _set_cell(sheet, r, c, None)

    for rate in rates:
        blank(summary)
        _set_cell(sheet, summary, 1, "[ 배관 부속재 ]")
        label = f"CD 전선관의 {round(rate * 100)} %" if abs(rate - CD_FITTING_RATE) < 1e-9 else f"전선관의 {round(rate * 100)} %"
        _set_cell(sheet, summary, 2, label)
        _set_cell(sheet, summary, 3, "식", align=CENTER)
        _set_cell(sheet, summary, 4, 1, align=RIGHT, number_format=QTY_FORMAT)
        c = letter(("rate", rate))
        _set_cell(sheet, summary, 5, f"=SUMPRODUCT(F{first}:F{last},{c}{first}:{c}{last})", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, summary, 6, f"=D{summary}*E{summary}*{rate}", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, summary, 11, f"=F{summary}+H{summary}+J{summary}", align=RIGHT, number_format=MONEY_FORMAT)
        material_rows.append(summary)
        summary += 1
    if any(ln.entry.has_sundry for ln in lines):
        blank(summary)
        _set_cell(sheet, summary, 1, "[ 소모 잡자재 ]")
        _set_cell(sheet, summary, 2, f"전선, 전선관의 {round(SUNDRY_RATE * 100)} %")
        _set_cell(sheet, summary, 3, "식", align=CENTER)
        _set_cell(sheet, summary, 4, 1, align=RIGHT, number_format=QTY_FORMAT)
        c = letter("sundry")
        _set_cell(sheet, summary, 5, f"=SUMPRODUCT(F{first}:F{last},{c}{first}:{c}{last})", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, summary, 6, f"=D{summary}*E{summary}*{SUNDRY_RATE}", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, summary, 11, f"=F{summary}+H{summary}+J{summary}", align=RIGHT, number_format=MONEY_FORMAT)
        material_rows.append(summary)
        summary += 1
    for job in jobs:
        blank(summary)
        _set_cell(sheet, summary, 1, "노 무 비")
        _set_cell(sheet, summary, 2, job)
        _set_cell(sheet, summary, 3, "인", align=CENTER)
        c = letter(("job", job))
        _set_cell(sheet, summary, 8, f"=SUMPRODUCT(D{first}:D{last},{c}{first}:{c}{last})", align=RIGHT, number_format=MONEY_FORMAT)
        labor_first = labor_first or summary
        labor_last = summary
        summary += 1
    if jobs:
        blank(summary)
        _set_cell(sheet, summary, 1, "[ 공 구 손 료 ]")
        _set_cell(sheet, summary, 2, f"노무비의 {round(TOOL_RATE * 100)} %")
        _set_cell(sheet, summary, 3, "식", align=CENTER)
        _set_cell(sheet, summary, 4, 1, align=RIGHT, number_format=QTY_FORMAT)
        _set_cell(sheet, summary, 7, f"=SUM(H{labor_first}:H{labor_last})", align=RIGHT, number_format=MONEY_FORMAT)
        _set_cell(sheet, summary, 8, f"=D{summary}*G{summary}*{TOOL_RATE}", align=RIGHT, number_format=MONEY_FORMAT)
        tool_row = summary
        summary += 1
    else:
        tool_row = None
    blank(summary)
    summary += 1
    total = summary
    blank(total)
    _set_cell(sheet, total, 1, "( 합       계 )")
    parts = [f"SUM(F{first}:F{last})"] + [f"F{r}" for r in material_rows]
    _set_cell(sheet, total, 6, "=" + "+".join(parts), align=RIGHT, number_format=MONEY_FORMAT)
    if tool_row:
        _set_cell(sheet, total, 8, f"=SUM(H{labor_first}:H{tool_row})", align=RIGHT, number_format=MONEY_FORMAT)
    _set_cell(sheet, total, 11, f"=F{total}+H{total}+J{total}", align=RIGHT, number_format=MONEY_FORMAT)

    _apply_sheet_look(sheet, total, LAST_COL, row_height=FORM_ROW_HEIGHT)
    for col, width in zip("ABCDEFGHIJKL", (36, 32, 7, 10, 12, 15, 13, 16, 10, 12, 16, 18)):
        sheet.column_dimensions[col].width = width
    for col in range(HELPER_START_COL, helper_last + 1):
        sheet.column_dimensions[get_column_letter(col)].hidden = True
    sheet.freeze_panes = "A5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def save_session(session: EstimateSession, sections: list[EstimateSection], dest_dir: Path | None) -> Path:
    """원본은 두고, 고른 공정마다 내역서 시트를 붙여 새 엑셀 파일로 저장한다."""
    workbook = session.workbook
    for old in session.estimate_sheets:  # 다시 저장할 때 내역서가 쌓이지 않게 먼저 지운다.
        if old in workbook.sheetnames:
            del workbook[old]
    session.estimate_sheets = []
    taken = set(workbook.sheetnames)
    names: list[str] = []
    for index, section in enumerate(sections):
        base = ESTIMATE_SHEET_NAME if index == 0 else f"{ESTIMATE_SHEET_NAME}({index + 1})"
        name = _sheet_name(base, taken)
        taken.add(name)
        names.append(name)
        write_estimate_sheet(workbook, name, section)
    # 일위대가목록 바로 뒤에 내역서를 둔다.
    anchor = workbook.sheetnames.index(ILWIDAE_LIST_SHEET_NAME) if ILWIDAE_LIST_SHEET_NAME in workbook.sheetnames else len(workbook.sheetnames) - 1
    for offset, name in enumerate(names, start=1):
        current = workbook.sheetnames.index(name)
        workbook.move_sheet(name, offset=anchor + offset - current)
    session.estimate_sheets = list(names)
    dest = build_result_path(dest_dir, mode="estimate")
    if session.source_path is not None:
        assert_safe_save(session.source_path, dest)
    workbook.save(dest)
    return dest
