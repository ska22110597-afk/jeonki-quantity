"""일위대가 호표 작성. 품명과 표준품셈 인부만 넣고, 부가세는 내역서 아래에 모은다."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.discipline import default_labor_name
from app.estimate_parse import normalize_header
from app.pumsam_aliases import alias_group
from app.excel_io import (
    apply_form_widths,
    AMOUNT_FORMAT,
    BODY_FONT,
    CENTER,
    FORM_ROW_HEIGHT,
    HEADER_FONT,
    ILWIDAE_SHEET_NAME,
    LEFT,
    PRICE_FORMAT,
    RIGHT,
    write_title_banner,
    _apply_sheet_look,
    _set_cell,
)
from app.fitting_rules import DEFAULT_FITTING_EXCLUDES, fitting_excluded
from app.items import LineItem
from app.pumsam import (
    PumsamRow,
    format_pumsam_ref,
    labor_kind_text,
    match_pumsam as match_pumsam_rows,
    pumsam_qty_value,
    pumsam_rate_value,
)
from app.wages import WAGES_SHEET_NAME, WageRow

PIPE_WASTE_QTY = 1.1
CD_FITTING_RATE = 0.40
CONDUIT_FITTING_RATE = 0.15
TELECOM_CONDUIT_FITTING_RATE = 0.20  # 통신 표준품셈 3-1-2: 박강·후강·합성수지·금속제 가요 전선관 20%
SUNDRY_RATE = 0.02
TOOL_RATE = 0.03
HO_TITLE_RE = re.compile(r"\(\s*호표\s*\d+\s*\)")


@dataclass
class IlwidaeBlock:
    item: LineItem
    ho_no: int
    title_row: int
    material_row: int
    sum_row: int
    labor_rows: list[int] = field(default_factory=list)
    pumsam_ref: str = ""
    fitting_rate: float = 0.0  # 전선관부속품비율(0 이면 부속품비 행 없음)
    has_sundry: bool = False  # 잡재료비 행 여부
    material_qty: float = 1.0  # 자재 수량(할증 포함)
    labor_jobs: list[str] = field(default_factory=list)
    labor_qtys: list[float] = field(default_factory=list)
    explain: list[list[Any]] = field(default_factory=list)  # 일위대가 품 해설 시트 행


_CONDUIT_ANCHORS = (
    "강제전선관",
    "경질비닐전선관",
    "CD전선관",
    "금속제가요전선관",
    "박강전선관",
    "합성수지파형관",
)
_WIRE_ANCHORS = ("옥내배선", "전력케이블", "제어용케이블", "UTP케이블", "광케이블")


def _in_anchor_group(name: Any, anchors: tuple[str, ...]) -> bool:
    group = alias_group(name)
    return any(anchor in group for anchor in anchors)


def is_conduit_name(name: Any) -> bool:
    """전선관과 같은 자재. HI관·후강관처럼 다른 말도 전선관으로 본다."""
    if _in_anchor_group(name, _CONDUIT_ANCHORS):
        return True
    text = str(name or "")
    return "전선관" in text and "부속" not in text


def is_cable_name(name: Any) -> bool:
    """전선·케이블. CV·HIV처럼 다른 말도 전선류로 본다. 전선관·트레이는 빼 둔다."""
    if is_conduit_name(name):
        return False
    if _in_anchor_group(name, _WIRE_ANCHORS):
        return True
    text = str(name or "")
    if any(word in text for word in ("트레이", "덕트", "몰딩")):
        return False
    return "케이블" in text or "전선" in text


def conduit_fitting_rate(name: Any, spec: Any = None, discipline: str = "전기") -> float:
    """전선관 부속품률. CD·PF(휨 합성수지관) 40%.

    전기 5-2: 후강·박강·합성수지 15%. 통신 3-1-2: 박강·후강·합성수지·금속제 가요 20%.
    """
    text = f"{name or ''} {spec or ''}".upper()
    if "CD" in text or "PF" in text or "주름" in text or ("합성수지" in text and "가요" in text):
        return CD_FITTING_RATE
    return TELECOM_CONDUIT_FITTING_RATE if discipline == "통신" else CONDUIT_FITTING_RATE


def is_meter_unit(unit: Any) -> bool:
    token = normalize_header(unit).lower()
    return token in {"m", "ｍ", "미터", "metre", "meter"}


def material_qty(item: LineItem, discipline: str = "전기") -> float:
    return material_waste_note(item, discipline)[0]


def material_waste_note(item: LineItem, discipline: str = "전기") -> tuple[float, str]:
    from app.material_waste import material_waste

    return material_waste(
        item.name,
        item.spec,
        item.unit,
        discipline,
        is_conduit=is_conduit_name(item.name),
        is_cable=is_cable_name(item.name),
    )


def match_pumsam(item: LineItem, rows: list[PumsamRow]) -> list[PumsamRow]:
    return match_pumsam_rows(item.name, item.spec, rows)


_PER_LENGTH = {"100m": 100.0, "10m": 10.0, "1000m": 1000.0, "km": 1000.0}
_PER_COUNT = {"10개": 10.0, "100개": 100.0, "10개소": 10.0, "100개소": 100.0}
_COUNT_UNITS = {"개", "ea", "개소", "pcs"}


def length_unit_divisor(row: PumsamRow, item_unit: Any) -> float:
    """품셈이 100 m·km 당으로 적혀 있고 단가대비표 품목이 m 단위면 몇으로 나눌지(아니면 1)."""
    book = re.sub(r"\s+", "", str(row.get("단위") or "")).lower()
    item = re.sub(r"\s+", "", str(item_unit or "")).lower()
    if item == "m":
        return _PER_LENGTH.get(book, 1.0)
    if item in _COUNT_UNITS:
        return _PER_COUNT.get(book, 1.0)
    return 1.0


def scale_to_unit(labors: list[PumsamRow], item_unit: Any) -> list[PumsamRow]:
    """100 m 당·km 당 품을 1 m 당으로 환산한 복사본을 돌려준다. 근거 칸에 환산 사실을 남긴다."""
    out: list[PumsamRow] = []
    for row in labors:
        divisor = length_unit_divisor(row, item_unit)
        if divisor == 1.0:
            out.append(row)
            continue
        copy = dict(row)
        copy["품셈"] = round(pumsam_qty_value(row) / divisor, 8)
        copy["환산"] = f"품셈 {str(row.get('단위')).strip()} 기준 → 1 {str(item_unit).strip()} 환산(÷{divisor:g})"
        out.append(copy)
    return out


def labor_qty_formula(row: PumsamRow) -> str:
    qty = pumsam_qty_value(row)
    rate = pumsam_rate_value(row)
    return f"={qty}*{rate}"


def labor_base_qty_formula(row: PumsamRow) -> str:
    """인부 수량 칸. 원표 품을 수식으로 둔다. 할증은 공량산출 할증%에서 한 번만 곱한다."""
    qty = pumsam_qty_value(row)
    text = f"{qty:.6f}".rstrip("0").rstrip(".")
    if text in {"", "-"}:
        text = "0"
    return f"={text}"


def applied_spec_note(item_spec: Any, row: PumsamRow) -> str:
    """단가대비표 규격과 다른 구간·규격의 품을 썼으면 어느 규격을 적용했는지 적는다."""
    from app.pumsam import spec_match_parts

    book_spec = str(row.get("규격") or "").strip()
    if not book_spec or not row.get("품셈근거"):
        return ""
    a_size, a_unit, _a_qual, _ = spec_match_parts(item_spec)
    b_size, b_unit, b_qual, _ = spec_match_parts(book_spec)
    if b_size is None:
        return ""
    if b_qual in ("이하", "이상", "미만", "초과"):
        return f"적용 규격 {book_spec}"
    if a_size is not None and a_unit == b_unit and abs(a_size - b_size) > 1e-9:
        return f"적용 규격 {book_spec}"
    return ""


def labor_qty_formula_with_place(row: PumsamRow) -> str:
    """인부 수량 칸. 노출·지중 같은 장소 할증이 붙은 품은 `=원표품*할증` 으로 보이게 둔다."""
    base = labor_base_qty_formula(row)
    rate = pumsam_rate_value(row)
    if abs(rate - 1.0) < 1e-9:
        return base
    return f"{base}*{rate:g}"


def quantity_ilwidae_refs(blocks: list[IlwidaeBlock]) -> list[tuple[str, int, list[int]]]:
    """공량산출서가 일위대가 호표를 찾을 때 쓰는 (키, 자재행, 인부행들)."""
    return [(block.item.key, block.material_row, list(block.labor_rows)) for block in blocks]


def renumber_ilwidae_ho(sheet: Worksheet) -> int:
    """시트에 적힌 호표를 위에서부터 1, 2, 3… 순서로 다시 붙인다."""
    count = 0
    for row in range(1, (sheet.max_row or 1) + 1):
        cell = sheet.cell(row, 1)
        value = cell.value
        if not isinstance(value, str) or "호표" not in value:
            continue
        count += 1
        label = f"( 호표 {count} )"
        if HO_TITLE_RE.search(value):
            cell.value = HO_TITLE_RE.sub(label, value, count=1)
        else:
            cell.value = re.sub(r"호표\s*\d+", f"호표 {count}", value, count=1)
    return count


def parse_ilwidae_blocks(sheet: Worksheet) -> list[IlwidaeBlock]:
    """이미 적힌 일위대가 시트에서 호표 구간을 읽는다."""
    blocks: list[IlwidaeBlock] = []
    title_row: int | None = None
    material_row: int | None = None
    labor_rows: list[int] = []
    name: Any = None
    spec: Any = None
    ho_no = 0

    def flush(sum_row: int | None = None) -> None:
        nonlocal title_row, material_row, labor_rows, name, spec, ho_no
        if title_row is None or material_row is None:
            title_row = None
            material_row = None
            labor_rows = []
            name = None
            spec = None
            return
        ho_no += 1
        blocks.append(
            IlwidaeBlock(
                item=LineItem(
                    excel_row=material_row,
                    name=name,
                    spec=spec,
                    unit=None,
                    qty=None,
                    material_price=None,
                ),
                ho_no=ho_no,
                title_row=title_row,
                material_row=material_row,
                sum_row=sum_row or (labor_rows[-1] + 1 if labor_rows else material_row + 1),
                labor_rows=list(labor_rows),
            )
        )
        title_row = None
        material_row = None
        labor_rows = []
        name = None
        spec = None

    for row in range(1, (sheet.max_row or 1) + 1):
        a = sheet.cell(row, 1).value
        b = sheet.cell(row, 2).value
        c = sheet.cell(row, 3).value
        a_text = str(a).strip() if a not in (None, "") else ""
        if not a_text and b in (None, ""):
            continue
        compact = a_text.replace(" ", "")
        if "호표" in a_text:
            if title_row is not None:
                flush()
            title_row = row
            continue
        if title_row is None:
            continue
        if "합계" in compact:
            flush(sum_row=row)
            continue
        unit = str(c).strip() if c not in (None, "") else ""
        if unit == "인":
            labor_rows.append(row)
            continue
        if material_row is None and a_text:
            material_row = row
            name = a
            spec = b
    if title_row is not None:
        flush()
    return blocks


NOTE_COL = 16  # 비고(13) 오른쪽 두 칸을 띄우고 적용 기준을 적는다.
_REF_RE = re.compile(r"^((?:전기|통신) ?\d+(?:-\d+)+)\s*(.*)$")


def split_pumsam_ref(text: str) -> tuple[str, str]:
    """`전기 4-31 해설⑧ …` → (`전기 4-31`, `해설⑧ …`). 비고 칸에는 앞부분만 둔다."""
    matched = _REF_RE.match(text or "")
    if not matched:
        return text or "", ""
    return matched.group(1), matched.group(2).strip()


EXPLAIN_SHEET_NAME = "일위대가 품 해설"
EXPLAIN_HEADERS = ["호표", "명칭", "규격", "단위", "적용 품셈(명칭 · 규격)", "품셈근거", "직종", "품셈 원표 품", "적용 할증", "일위대가 인부 수량", "자재 부속·잡재료", "적용 방법 해설"]


def _explain_rows(ho_no, item, labors, found, ref, ref_extra, qty, waste_note, fitting_rate, has_sundry):
    rows: list[list[Any]] = []
    extras: list[str] = []
    if fitting_rate:
        label = "CD전선관" if abs(fitting_rate - CD_FITTING_RATE) < 1e-9 else "전선관"
        extras.append(f"전선관부속품비 {label}의 {round(fitting_rate * 100)}%")
    if has_sundry:
        extras.append(f"잡재료비 {round(SUNDRY_RATE * 100)}%")
    if qty and abs(qty - 1) > 1e-9:
        extras.append(f"자재 할증 반영 수량 {qty:g}" + (f" ({waste_note})" if waste_note else ""))
    for index, labor in enumerate(labors):
        if not found:
            how = "품셈표에 없어 인부 수량 0 — " + no_pumsam_reason(item.name, item.spec)
        else:
            parts: list[str] = []
            book = str(labor.get("명칭") or "")
            lspec = str(labor.get("규격") or "")
            if book and lspec and item.spec and lspec.replace(" ", "") != str(item.spec).replace(" ", ""):
                parts.append(f"단가대비표 규격과 달라 품셈 '{lspec}' 구간 적용")
            rate = pumsam_rate_value(labor)
            if abs(rate - 1) > 1e-9:
                parts.append(f"장소·조건 할증 {round(rate * 100)}% 곱함")
            if labor.get("환산"):
                parts.append(str(labor["환산"]))
            if "CD" in f"{book}{lspec}".upper() or "PE" in lspec.upper():
                parts.append("CD·PE관: 합성수지 전선관 품의 80% (100mm 이상은 100%) — 품셈표에 환산해 둔 값")
            if ref_extra:
                parts.append(ref_extra)
            how = " · ".join(parts) or "품셈표 값을 그대로 적용"
        rows.append([
            ho_no if index == 0 else None,
            item.name if index == 0 else None,
            item.spec if index == 0 else None,
            item.unit if index == 0 else None,
            f"{labor.get('명칭') or ''} · {labor.get('규격') or ''}".strip(" ·") if found else "-",
            ref if (found and index == 0) else None,
            labor.get("노무명칭"),
            labor.get("품셈") if found else 0,
            f"{round(pumsam_rate_value(labor) * 100)}%" if found else None,
            pumsam_qty_value(labor) * pumsam_rate_value(labor),
            ", ".join(extras) if index == 0 else None,
            how,
        ])
    return rows


def write_explain_sheet(sheet: Worksheet, blocks: list["IlwidaeBlock"]) -> None:
    """일위대가마다 품셈을 어떻게 적용했는지 한 줄씩 적는다."""
    sheet.title = EXPLAIN_SHEET_NAME
    sheet.append(EXPLAIN_HEADERS)
    for block in blocks:
        for row in block.explain:
            sheet.append(row)
    widths = (7, 28, 24, 6, 34, 11, 12, 11, 9, 13, 34, 70)
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        row[11].alignment = Alignment(wrap_text=True, vertical="top")
        row[7].number_format = "General"
        row[9].number_format = "General"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(sheet.max_row, 1)}"


_NO_PUMSAM_ACCESSORY = ("부속", "부품", "행거", "커넥터", "앵커", "볼트", "너트", "브라켓", "BRACKET", "커버", "COVER", "클램프", "CLAMP", "새들", "철물", "서포트", "고정", "밴드", "엘보")


def _note(sheet: Worksheet, row: int, text: str | None) -> None:
    """적용 기준 칸. 표 테두리 밖 빈 칸에 글자만 둔다."""
    cell = sheet.cell(row=row, column=NOTE_COL, value=text or None)
    cell.font = BODY_FONT
    cell.alignment = LEFT


_PANEL_WORDS = ("분전반", "배전반", "저압반", "동력반", "변압기반", "특고압반", "제어반", "차단기함", "개폐기", "발전기", "수배전")
_CIVIL_WORDS = ("터파기", "되메우기", "맨홀", "레미콘", "철근", "합판", "각재", "형틀", "이형제", "WIRE MESH", "장비운반", "기계경비", "코아", "비트", "리밍", "천공", "부가가치세", "방화포", "조명탑")
_MATERIAL_WORDS = ("와샤", "WASHER", "너트", "NUT", "볼트", "BOLT", "철선", "철못", "산소", "아세틸렌", "철판", "강판", "강봉", "환봉", "형광", "용융아연", "RAIL", "CHANNEL", "JOINER", "MOLD", "POWDER", "FLINT", "표지시트", "애자", "밴드", "단자카바", "스리브", "슬리브", "접지크램프", "저감제", "수막처리")
_TELECOM_WORDS = ("광케이블", "UTP", "MODULE", "CPU", "POWER SUPPLY", "I/O", "통신", "감시장치", "프로그래밍", "연동시험", "종합시험", "원격")


def small_cable_joint(name: Any, spec: Any = None) -> bool:
    """10㎟ 이하 전력케이블 단말처리·직선접속. 4-36/4-37 해설④ 에 따라 5-13 포설 품에 포함되므로 따로 품을 붙이지 않는다."""
    from app.pumsam import spec_match_parts

    text = f"{name or ''}"
    if not (("단말처리" in text or "직선접속" in text) and "케이블" in text):
        return False
    size, unit, _qual, _ = spec_match_parts(spec)
    return size is not None and str(unit or "mm2") in ("mm2", "㎟") and size < 16


def no_pumsam_reason(name: Any, spec: Any = None) -> str:
    """품셈 수량이 0으로 들어간 이유(비고 오른쪽 칸에 적는다)."""
    text = f"{name or ''} {spec or ''}".upper()
    if any(w in text for w in ("단말처리", "직선접속")) and "케이블" in text:
        return "품셈 없음: 10㎟ 이하 케이블 단말·접속은 4-36/4-37 해설④에 따라 5-13 포설 품에 포함(별도 계상 안 함)"
    if any(word.upper() in text for word in _NO_PUMSAM_ACCESSORY):
        return "품셈 없음: 부속품(모체 자재 설치품에 포함)"
    if any(word.upper() in text for word in _CIVIL_WORDS):
        return "품셈 없음: 전기 품셈 범위 밖(토목·건축·장비 공종)"
    if any(word.upper() in text for word in _TELECOM_WORDS):
        return "품셈 없음: 통신·제어 품셈 대상(통신 파트로 생성하거나 직접 입력)"
    if any(word.upper() in text for word in _PANEL_WORDS):
        return "품셈 없음: 규격에 용량(AF·회로수·kVA)이 없어 품셈을 고를 수 없음(품셈 연결로 지정)"
    if any(word.upper() in text for word in _MATERIAL_WORDS):
        return "품셈 없음: 자재 단품(별도 설치 품셈 없음)"
    return "품셈 없음: 해당 품셈을 찾지 못함(직접 입력)"


def _write_ilwidae_header(sheet: Worksheet) -> None:
    write_title_banner(sheet, "일 위 대 가", 13)
    labels = {
        1: "품      명",
        2: "규      격",
        3: "단위",
        4: "수량",
        5: "재  료  비",
        7: "노  무  비",
        9: "경      비",
        11: "합      계",
        13: "비  고",
    }
    for col in range(1, 14):
        _set_cell(sheet, 3, col, labels.get(col), font=HEADER_FONT, align=CENTER)
        _set_cell(sheet, 4, col, None, font=HEADER_FONT, align=CENTER)
    for col, text in (
        (5, "단  가"),
        (6, "금  액"),
        (7, "단  가"),
        (8, "금  액"),
        (9, "단  가"),
        (10, "금  액"),
        (11, "단  가"),
        (12, "금  액"),
    ):
        _set_cell(sheet, 4, col, text, font=HEADER_FONT, align=CENTER)
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    sheet.merge_cells("D3:D4")
    sheet.merge_cells("E3:F3")
    sheet.merge_cells("G3:H3")
    sheet.merge_cells("I3:J3")
    sheet.merge_cells("K3:L3")
    sheet.merge_cells("M3:M4")


def _price(sheet: Worksheet, row: int, col: int, value: Any, align=RIGHT) -> None:
    _set_cell(sheet, row, col, value, font=BODY_FONT, align=align, number_format=PRICE_FORMAT)


def _amount(sheet: Worksheet, row: int, col: int, value: Any, align=RIGHT) -> None:
    _set_cell(sheet, row, col, value, font=BODY_FONT, align=align, number_format=AMOUNT_FORMAT)


def _cost_totals(sheet: Worksheet, row: int) -> None:
    """단가=재료+노무+경비(소수 2자리), 금액=재료+노무+경비(소수 1자리)."""
    _price(sheet, row, 11, f"=TRUNC(E{row}+G{row}+I{row},2)")
    _amount(sheet, row, 12, f"=TRUNC(F{row}+H{row}+J{row},1)")


def _zero_price() -> str:
    return "=0"


def _qty_times_price(row: int, price_col: str) -> str:
    """금액 = 단가×수량, 소수 1자리."""
    return f"=TRUNC(D{row}*{price_col}{row},1)"


def _qty_times_price_trunc(row: int, price_col: str) -> str:
    """노무·경비 금액. 소수 1자리."""
    return f"=TRUNC({price_col}{row}*D{row},1)"


def _idle_labor_expense(sheet: Worksheet, row: int) -> None:
    """재료 행에서 쓰지 않는 노무·경비 칸도 수식으로 둔다."""
    _price(sheet, row, 7, _zero_price())
    _amount(sheet, row, 8, _qty_times_price_trunc(row, "G"))
    _price(sheet, row, 9, _zero_price())
    _amount(sheet, row, 10, _qty_times_price_trunc(row, "I"))


def _idle_material_expense(sheet: Worksheet, row: int) -> None:
    """노무 행에서 쓰지 않는 재료·경비 칸도 수식으로 둔다."""
    _price(sheet, row, 5, _zero_price())
    _amount(sheet, row, 6, _qty_times_price(row, "E"))
    _price(sheet, row, 9, _zero_price())
    _amount(sheet, row, 10, _qty_times_price_trunc(row, "I"))


def _price_ref(item: LineItem, compare_sheet: str) -> str:
    col_index = (item.price_col if item.price_col is not None else 11) + 1
    return f"='{compare_sheet}'!{get_column_letter(col_index)}{item.excel_row}"


def _write_percent_material_row(
    sheet: Worksheet,
    row: int,
    name: str,
    spec: str,
    material_row: int,
    rate: float,
) -> None:
    """부속품·잡재료. 규격에 비율을 적고, 수식은 재료비 금액(F)에 둔다."""
    _set_cell(sheet, row, 1, name, font=BODY_FONT)
    _set_cell(sheet, row, 2, spec, font=BODY_FONT)
    _set_cell(sheet, row, 3, "식", font=BODY_FONT, align=CENTER)
    _set_cell(sheet, row, 4, 1, font=BODY_FONT, align=RIGHT, number_format="General")
    _price(sheet, row, 5, f"=IF(D{row}=0,0,TRUNC(F{row}/D{row},2))")
    _amount(sheet, row, 6, f"=TRUNC(F{material_row}*{rate},1)")
    _idle_labor_expense(sheet, row)
    _cost_totals(sheet, row)
    _set_cell(sheet, row, 13, None)


def _write_tool_loss_row(sheet: Worksheet, row: int, labor_rows: list[int]) -> None:
    """기본 행. 직접노무비의 3%를 경비 금액에 넣는다."""
    _set_cell(sheet, row, 1, "공구손료", font=BODY_FONT)
    _set_cell(sheet, row, 2, "직접노무비의 3%", font=BODY_FONT)
    _set_cell(sheet, row, 3, "식", font=BODY_FONT, align=CENTER)
    _set_cell(sheet, row, 4, 1, font=BODY_FONT, align=RIGHT, number_format="General")
    _price(sheet, row, 5, _zero_price())
    _amount(sheet, row, 6, _qty_times_price(row, "E"))
    _price(sheet, row, 7, _zero_price())
    _amount(sheet, row, 8, _qty_times_price_trunc(row, "G"))
    if labor_rows:
        parts = "+".join(f"H{labor}" for labor in labor_rows)
        tool_amount = f"=TRUNC(({parts})*{TOOL_RATE},1)"
    else:
        tool_amount = "=0"
    _price(sheet, row, 9, f"=IF(D{row}=0,0,TRUNC(J{row}/D{row},2))")
    _amount(sheet, row, 10, tool_amount)
    _cost_totals(sheet, row)
    _set_cell(sheet, row, 13, None)


def write_ilwidae_sheet(
    sheet: Worksheet,
    items: list[LineItem],
    pumsam_rows: list[PumsamRow],
    wage_rows: list[WageRow],
    compare_sheet: str = "단가대비표",
    discipline: str | None = None,
    fitting_excludes: tuple[str, ...] = DEFAULT_FITTING_EXCLUDES,
) -> list[IlwidaeBlock]:
    sheet.title = ILWIDAE_SHEET_NAME
    _write_ilwidae_header(sheet)
    cursor = 5
    blocks: list[IlwidaeBlock] = []
    ho_no = 0
    fallback_labor = default_labor_name(discipline)
    seen_keys: set[str] = set()
    for item in items:
        if item.section or (item.spec is None and item.unit is None):
            continue
        if item.key in seen_keys:  # 같은 품명·규격은 호표를 한 번만 만든다.
            continue
        seen_keys.add(item.key)
        ho_no += 1
        labors = [] if small_cable_joint(item.name, item.spec) else match_pumsam_rows(item.name, item.spec, pumsam_rows)
        pumsam_found = bool(labors)
        labors = scale_to_unit(labors, item.unit)
        if not labors:
            labors = [
                {
                    "명칭": item.name,
                    "규격": item.spec,
                    "노무명칭": fallback_labor,
                    "품셈": 0,
                    "할증%": 100,
                    "품셈근거": "",
                }
            ]
        ref, ref_extra = split_pumsam_ref(format_pumsam_ref(labors[0].get("품셈근거") or ""))
        title = f"{item.name} {item.spec or ''}  ( 호표 {ho_no} )".strip()
        title_row = cursor
        _set_cell(sheet, cursor, 1, title, font=BODY_FONT, align=LEFT)
        for col in range(2, 14):
            _set_cell(sheet, cursor, col, None)
        sheet.merge_cells(start_row=cursor, start_column=1, end_row=cursor, end_column=13)
        cursor += 1

        material_row = cursor
        qty, waste_note = material_waste_note(item, discipline or "전기")
        _set_cell(sheet, cursor, 1, item.name, font=BODY_FONT)
        _set_cell(sheet, cursor, 2, item.spec, font=BODY_FONT)
        _set_cell(sheet, cursor, 3, item.unit or "개", font=BODY_FONT, align=CENTER)
        _set_cell(sheet, cursor, 4, qty, font=BODY_FONT, align=RIGHT, number_format="General")
        _price(sheet, cursor, 5, _price_ref(item, compare_sheet))
        _amount(sheet, cursor, 6, _qty_times_price(cursor, "E"))
        _idle_labor_expense(sheet, cursor)
        _cost_totals(sheet, cursor)
        _set_cell(sheet, cursor, 13, ref or None, font=BODY_FONT, align=LEFT)
        _note(sheet, cursor, waste_note)
        cursor += 1

        conduit = is_conduit_name(item.name)
        wire = is_cable_name(item.name)
        fitting_rate = 0.0
        if conduit and not fitting_excluded(item.name, item.spec, fitting_excludes):
            fitting_rate = conduit_fitting_rate(item.name, item.spec, discipline or "전기")
            _write_percent_material_row(
                sheet,
                cursor,
                "전선관부속품비",
                (
                    f"CD전선관의 {round(fitting_rate * 100)}%"
                    if abs(fitting_rate - CD_FITTING_RATE) < 1e-9
                    else f"전선관의 {round(fitting_rate * 100)}%"
                ),
                material_row,
                fitting_rate,
            )
            cursor += 1
        if conduit or wire:
            sundry_spec = "배관의 2%" if conduit else "배선의 2%"
            _write_percent_material_row(
                sheet, cursor, "잡재료비", sundry_spec, material_row, SUNDRY_RATE
            )
            cursor += 1

        labor_rows: list[int] = []
        for labor_index, labor in enumerate(labors):
            job = labor.get("노무명칭") or fallback_labor
            _set_cell(sheet, cursor, 1, job, font=BODY_FONT)
            _set_cell(sheet, cursor, 2, labor_kind_text(labor), font=BODY_FONT)
            _set_cell(sheet, cursor, 3, "인", font=BODY_FONT, align=CENTER)
            _set_cell(
                sheet,
                cursor,
                4,
                labor_qty_formula_with_place(labor),
                font=BODY_FONT,
                align=RIGHT,
                number_format="General",
            )
            job_lit = str(job).replace('"', '""')
            _idle_material_expense(sheet, cursor)
            _price(
                sheet,
                cursor,
                7,
                f"=IFERROR(VLOOKUP(\"{job_lit}\",'{WAGES_SHEET_NAME}'!A:B,2,FALSE),0)",
            )
            _amount(sheet, cursor, 8, f"=TRUNC(G{cursor}*D{cursor},1)")
            _cost_totals(sheet, cursor)
            if labor_index == 0:
                notes: list[str] = []
                rate_pct = round(pumsam_rate_value(labor) * 100)
                if rate_pct != 100:
                    notes.append(f"장소 품 할증 {rate_pct}%")
                if labor.get("환산"):
                    notes.append(str(labor["환산"]))
                applied = applied_spec_note(item.spec, labor)
                if applied:
                    notes.append(applied)
                if ref_extra:
                    notes.append(ref_extra)
                if pumsam_found and ("CD" in f"{labor.get('명칭')}{labor.get('규격')}".upper() or "PE" in str(labor.get("규격") or "").upper()):
                    notes.append("합성수지 전선관 품의 80%(100mm 이상 100%)")
                if not pumsam_found:
                    notes.append(no_pumsam_reason(item.name, item.spec))
                _set_cell(sheet, cursor, 13, ref or None, font=BODY_FONT, align=LEFT)
                _note(sheet, cursor, " · ".join(notes))
            else:
                _set_cell(sheet, cursor, 13, None)
            labor_rows.append(cursor)
            cursor += 1

        _write_tool_loss_row(sheet, cursor, labor_rows)
        cursor += 1

        first_data = material_row
        last_data = cursor - 1
        sum_row = cursor
        _set_cell(sheet, cursor, 1, " [ 합          계 ]", font=BODY_FONT)
        for col in range(2, 14):
            _set_cell(sheet, cursor, col, None)
        _amount(sheet, cursor, 6, f"=SUM(F{first_data}:F{last_data})")
        _amount(sheet, cursor, 8, f"=SUM(H{first_data}:H{last_data})")
        _amount(sheet, cursor, 10, f"=SUM(J{first_data}:J{last_data})")
        _amount(sheet, cursor, 12, f"=TRUNC(F{cursor}+H{cursor}+J{cursor},1)")
        _set_cell(sheet, cursor, 13, None)
        explain_rows = _explain_rows(
            ho_no, item, labors, pumsam_found, ref, ref_extra, qty, waste_note,
            fitting_rate, bool(conduit or wire),
        )
        blocks.append(
            IlwidaeBlock(
                item=item,
                ho_no=ho_no,
                title_row=title_row,
                material_row=material_row,
                sum_row=sum_row,
                labor_rows=labor_rows,
                pumsam_ref=ref,
                fitting_rate=fitting_rate,
                has_sundry=bool(conduit or wire),
                material_qty=qty,
                labor_jobs=[str(l.get("노무명칭") or fallback_labor) for l in labors],
                labor_qtys=[pumsam_qty_value(l) * pumsam_rate_value(l) for l in labors],
                explain=explain_rows,
            )
        )
        cursor += 2

    _apply_sheet_look(sheet, max(cursor, 4), 13, row_height=FORM_ROW_HEIGHT)
    apply_form_widths(sheet, 13, remark_col=13)
    return blocks
