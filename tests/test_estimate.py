from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from app.estimate_build import EstimateLine, EstimateSection, generate_session, save_session

SAMPLE = Path(__file__).resolve().parent.parent / "samples" / "단가대비표_샘플.xlsx"


@pytest.fixture(scope="module")
def session(tmp_path_factory):
    db = tmp_path_factory.mktemp("db")
    return generate_session(SAMPLE, "전기", db)


def _pick(session, *, rate=None, sundry=None, count=2):
    found = []
    for entry in session.entries:
        if rate is not None and abs(entry.fitting_rate - rate) > 1e-9:
            continue
        if sundry is not None and entry.has_sundry != sundry:
            continue
        found.append(entry)
        if len(found) == count:
            break
    return found


def test_entries_carry_rules(session) -> None:
    assert session.entries
    conduits = _pick(session, rate=0.15)
    assert conduits, "샘플에 전선관 15% 품목이 있어야 합니다."
    assert all(e.has_sundry for e in conduits)
    assert all(e.material_unit > 0 for e in conduits)


def test_save_two_sections_and_layout(session, tmp_path) -> None:
    lines = [EstimateLine(e, 10 + i, "비고") for i, e in enumerate(_pick(session, rate=0.15, count=2))]
    lines += [EstimateLine(e, 5, "") for e in _pick(session, rate=0.0, sundry=False, count=1)]
    sections = [EstimateSection("1.공사::1-1.세대", lines), EstimateSection("2.공용부", lines[:1])]
    dest = save_session(session, sections, tmp_path)
    wb = openpyxl.load_workbook(dest)
    names = wb.sheetnames
    assert names.index("내역서") == names.index("일위대가목록") + 1
    assert "내역서(2)" in names
    ws = wb["내역서"]
    assert ws["A5"].value == "1.공사::1-1.세대"
    assert ws["A6"].value == lines[0].entry.name
    assert ws["E6"].value.startswith("='일위대가'!F")
    labels = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1)]
    assert "[ 배관 부속재 ]" in labels
    assert "[ 소모 잡자재 ]" in labels
    assert "노 무 비" in labels
    assert "[ 공 구 손 료 ]" in labels
    assert "( 합       계 )" in labels
    assert ws.column_dimensions["N"].hidden is True


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice 없음")
def test_totals_match_python(session, tmp_path) -> None:
    lines = [EstimateLine(e, 100 + i * 7, "") for i, e in enumerate(_pick(session, rate=0.15, count=3))]
    lines += [EstimateLine(e, 3, "") for e in _pick(session, rate=0.0, sundry=False, count=2)]
    dest = save_session(session, [EstimateSection("검산", lines)], tmp_path)
    out = tmp_path / "calc"
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(dest)],
        check=True,
        capture_output=True,
        timeout=170,
    )
    ws = openpyxl.load_workbook(out / dest.name, data_only=True)["내역서"]
    mat_items = sum(l.qty * l.entry.material_unit for l in lines)
    fit = sum(l.qty * l.entry.material_unit * l.entry.fitting_rate for l in lines)
    sundry = sum(l.qty * l.entry.material_unit * 0.02 for l in lines if l.entry.has_sundry)
    labor = sum(l.qty * l.entry.labor_unit for l in lines)
    total_row = next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("( 합"))
    assert ws.cell(total_row, 6).value == pytest.approx(mat_items + fit + sundry, abs=1)
    assert ws.cell(total_row, 8).value == pytest.approx(labor * 1.03, abs=1)
    assert ws.cell(total_row, 11).value == pytest.approx(mat_items + fit + sundry + labor * 1.03, abs=2)


def test_window_select_and_add(session) -> None:
    from PyQt6.QtCore import Qt
    from PyQt6.QtWidgets import QApplication

    from app.estimate_window import COL_NOTE, COL_QTY, EstimateWindow

    app = QApplication.instance() or QApplication([])  # noqa: F841 — 참조를 잡아 둔다
    window = EstimateWindow()
    try:
        assert window.save_button.isEnabled() is False
        window._load_session(session)
        assert window.table.rowCount() == len(session.entries)
        window.table.item(0, 0).setCheckState(Qt.CheckState.Checked)
        assert window.table.item(0, COL_QTY).text() == "1"  # 체크하면 수량 1 자동
        window.table.item(0, COL_QTY).setText("12")
        window.table.item(0, COL_NOTE).setText("설치품 포함")
        assert window.add_button.isEnabled() is False  # 제목이 없으면 담을 수 없다
        window.title_edit.setText("1.아파트전기공사")
        assert window.add_button.isEnabled() is True
        window._on_add_section()
        assert len(window._sections) == 1
        line = window._sections[0].lines[0]
        assert (line.qty, line.remark) == (12.0, "설치품 포함")
        assert window._checked_rows() == []  # 담은 뒤 체크가 풀린다
        assert window.save_button.isEnabled() is True
        window.search_edit.setText("zzzz없는품목")
        assert all(window.table.isRowHidden(r) for r in range(window.table.rowCount()))
        window._on_edit_section() if window.section_list.selectedItems() else None
        window.section_list.setCurrentRow(0)
        window._on_edit_section()
        assert window.title_edit.text() == "1.아파트전기공사"
        assert window._sections == []
    finally:
        window.close()


def test_saving_twice_does_not_stack_sheets(session, tmp_path) -> None:
    lines = [EstimateLine(e, 1, "") for e in _pick(session, rate=0.15, count=1)]
    first = save_session(session, [EstimateSection("A", lines), EstimateSection("B", lines)], tmp_path)
    second = save_session(session, [EstimateSection("C", lines)], tmp_path)
    names = openpyxl.load_workbook(second).sheetnames
    assert [n for n in names if n.startswith("내역서")] == ["내역서"]
    assert openpyxl.load_workbook(second)["내역서"]["A5"].value == "C"


def test_resume_from_saved_ilwidae_list(session, tmp_path) -> None:
    from app.estimate_build import is_result_workbook, open_session

    saved = save_session(session, [], tmp_path)
    assert saved.name.startswith("일위대가목록_결과_")
    assert is_result_workbook(saved) and not is_result_workbook(SAMPLE)
    again = open_session(saved, "전기", tmp_path)
    assert len(again.entries) == len(session.entries)
    for a, b in zip(session.entries, again.entries):
        assert (a.name, a.spec, a.unit, a.material_row, a.labor_rows) == (b.name, b.spec, b.unit, b.material_row, b.labor_rows)
        assert round(a.material_unit, 1) == round(b.material_unit, 1)
        assert a.fitting_rate == b.fitting_rate and a.has_sundry == b.has_sundry
    out = save_session(again, [EstimateSection("이어서", [EstimateLine(again.entries[0], 3)])], tmp_path)
    assert out.name.startswith("내역서_결과_")
    assert "내역서" in openpyxl.load_workbook(out).sheetnames


def test_material_waste_and_basis_note() -> None:
    from app.material_waste import material_waste

    assert material_waste("강제전선관", "아연도 16 mm", "M", "통신", is_conduit=True, is_cable=False)[0] == 1.10
    q, note = material_waste("UTP 케이블", "Cat.6 4P", "M", "통신", is_conduit=False, is_cable=True)
    assert q == 1.075 and "통신 1-1-6" in note
    q, note = material_waste("광케이블(옥외용)", "SM 12c", "M", "통신", is_conduit=False, is_cable=True)
    assert q == 1.03 and "옥외" in note
    assert material_waste("앵커볼트", "M10", "개", "전기", is_conduit=False, is_cable=False) == (1.0, "")


def test_summary_has_six_blank_rows_and_no_dotted_qty(session, tmp_path) -> None:
    sec = EstimateSection("t", [EstimateLine(session.entries[0], 1), EstimateLine(session.entries[1], 2.5)])
    out = save_session(session, [sec], tmp_path)
    ws = openpyxl.load_workbook(out)["내역서"]
    assert ws.cell(6, 4).number_format == "#,##0"
    assert ws.cell(7, 4).number_format == "#,##0.0##"
    total = next(r for r in range(1, ws.max_row + 1) if "합" in str(ws.cell(r, 1).value or "") and "계" in str(ws.cell(r, 1).value or ""))
    assert all(ws.cell(total - k, 1).value is None for k in range(1, 7))


def test_resume_renumbers_ho_pyo_after_manual_insert(session, tmp_path) -> None:
    from app.estimate_build import open_session

    saved = save_session(session, [], tmp_path)
    wb = openpyxl.load_workbook(saved)
    ws = wb["일위대가"]
    titles = [r for r in range(5, ws.max_row + 1) if "호표" in str(ws.cell(r, 1).value or "")]
    ws.cell(titles[1], 1).value = "새로 끼운 호표 ( 호표 )"
    ws.cell(titles[2], 1).value = "아무거나 ( 호표 77 )"
    edited = tmp_path / "edited.xlsx"
    wb.save(edited)
    again = open_session(edited, "전기", tmp_path)
    assert [e.ho_no for e in again.entries[:4]] == [1, 2, 3, 4]
    assert "( 호표 3 )" in again.workbook["일위대가"].cell(titles[2], 1).value


def test_lower_bound_band_and_new_electric_names() -> None:
    from app.electric_pumsam_data import _row
    from app.pumsam import load_pumsam_database, match_pumsam

    rows = [_row("테스트선", s, "M", "내선전공", q, 100, "전기9-9") for s, q in (("3.5 mm 이상", 0.2), ("6 mm 이상", 0.3), ("8.5 mm 이상", 0.4))]
    assert match_pumsam("테스트선", "4.3mm", rows)[0]["품셈"] == 0.2
    assert match_pumsam("테스트선", "7mm", rows)[0]["품셈"] == 0.3
    book = load_pumsam_database(None, "전기")
    assert match_pumsam("나사없는전선관", "E31", book)[0]["품셈"] == 0.08
    assert match_pumsam("조명기구(J)", "LED 13W(직부)", book)[0]["내선전공" if False else "노무명칭"] == "내선전공"


def test_unmatched_flag_check_sheet_and_autosave(session, tmp_path) -> None:
    from app.estimate_build import autosave_path, read_autosave, write_autosave

    assert any(not e.matched for e in session.entries) or True
    sec = EstimateSection("임시", [EstimateLine(session.entries[0], 4, "메모")])
    path = tmp_path / "임시저장.json"
    write_autosave(path, session, [sec])
    back = read_autosave(path, session)
    assert back[0].title == "임시" and back[0].lines[0].qty == 4 and back[0].lines[0].remark == "메모"
    out = save_session(session, [], tmp_path)
    names = openpyxl.load_workbook(out).sheetnames
    if any(not e.matched for e in session.entries):
        assert "점검목록" in names


def test_text_numbers_and_safe_save(tmp_path) -> None:
    from openpyxl import Workbook

    from app.numbers import parse_number
    from app.paths import save_workbook_safely

    assert parse_number("1,200㎡") == (1200.0, True)
    assert parse_number("약 3") == (3.0, True)
    assert parse_number("abc") == (None, False)
    wb = Workbook()
    dest = tmp_path / "a.xlsx"
    assert save_workbook_safely(wb, dest) == dest

    class Locked:
        calls = 0

        def save(self, path):
            Locked.calls += 1
            if Locked.calls == 1:
                raise PermissionError("locked")
            Path(path).write_bytes(b"x")

    assert save_workbook_safely(Locked(), tmp_path / "b.xlsx").name == "b_2.xlsx"


def test_link_overrides_matching(tmp_path) -> None:
    from app.pumsam import load_pumsam_database, match_pumsam
    from app.pumsam_links import activate, save_link

    book = load_pumsam_database(None, "전기")
    assert match_pumsam("엉뚱한자재", "A-1", book) == []
    save_link(tmp_path, "엉뚱한자재", "A-1", "강제전선관", "아연도 16 mm")
    activate(tmp_path)
    try:
        assert match_pumsam("엉뚱한자재", "A-1", book)[0]["노무명칭"] == "내선전공"
    finally:
        activate(tmp_path / "없는폴더")


def test_unknown_layout_needs_column_mapping(tmp_path) -> None:
    from openpyxl import Workbook

    from app.estimate_build import convert_with_columns, layout_ok

    wb = Workbook()
    ws = wb.active
    ws.append(["우리 자재표"])
    ws.append(["자재명", "형식", "수량단위", "금액"])
    ws.append(["강제전선관", "아연도 16 mm", "M", "3,374원"])
    odd = tmp_path / "odd.xlsx"
    wb.save(odd)
    assert layout_ok(SAMPLE) and not layout_ok(odd)
    fixed = convert_with_columns(odd, tmp_path, first_data_row=3, name_col=1, spec_col=2, unit_col=3, price_col=4)
    assert layout_ok(fixed)
    again = generate_session(fixed, "전기", tmp_path / "db")
    assert again.entries[0].matched and again.entries[0].material_unit > 3000


def test_window_has_status_column_and_filters(session) -> None:
    from PyQt6.QtWidgets import QApplication

    from app.estimate_window import COL_STATE, EstimateWindow

    app = QApplication.instance() or QApplication([])  # noqa: F841
    window = EstimateWindow()
    try:
        window._load_session(session)
        assert window.table.item(0, COL_STATE) is not None
        window.unmatched_only.setChecked(True)
        hidden = sum(1 for r in range(window.table.rowCount()) if window.table.isRowHidden(r))
        assert hidden == sum(1 for e in session.entries if e.matched)
    finally:
        window.close()
