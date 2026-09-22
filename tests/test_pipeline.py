from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.excel_io import (
    COMPARE_SHEET_NAME,
    ESTIMATE_SHEET_NAME,
    ILWIDAE_LIST_SHEET_NAME,
    ILWIDAE_SHEET_NAME,
    QUANTITY_SHEET_NAME,
    save_result_workbook,
)
from app.ilwidae import labor_qty_formula, match_pumsam
from app.items import LineItem
from app.pumsam import PUMSAM_SHEET_NAME, default_pumsam_rows
from app.wages import WAGES_SHEET_NAME, default_wage_rows


def _write_compare(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "단가대비표"
    sheet["A1"] = "단 가 대 비 표"
    sheet["A2"] = "명칭"
    sheet["B2"] = "규격"
    sheet["C2"] = "단위"
    sheet["D2"] = "수량"
    sheet["E2"] = "재료비"
    sheet["E3"] = "단가"
    sheet["F3"] = "금액"
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet.merge_cells("C2:C3")
    sheet.merge_cells("D2:D3")
    sheet.merge_cells("E2:F2")
    sheet["A4"] = "경질비닐전선관_지중"
    sheet["B4"] = "HI 16 mm"
    sheet["C4"] = "M"
    sheet["D4"] = 100
    sheet["E4"] = 200
    sheet["A5"] = "경질비닐전선관_노출"
    sheet["B5"] = "HI 104 mm"
    sheet["C5"] = "M"
    sheet["D5"] = 40
    sheet["E5"] = 151
    sheet["A6"] = "배선용단자함"
    sheet["B6"] = "10 P 이하"
    sheet["C6"] = "대"
    sheet["D6"] = 1
    sheet["E6"] = 50000
    workbook.save(path)
    workbook.close()


def test_conduit_extras_and_two_labors(tmp_path: Path) -> None:
    source = tmp_path / "단가대비표.xlsx"
    _write_compare(source)
    dest = save_result_workbook(unit_price_path=source, dest_dir=tmp_path / "out")
    result = load_workbook(dest, data_only=False)
    try:
        assert dest.name.startswith("일위대가목록_결과_")
        assert result.sheetnames[0] == COMPARE_SHEET_NAME
        assert ILWIDAE_SHEET_NAME in result.sheetnames
        assert ESTIMATE_SHEET_NAME not in result.sheetnames
        assert ILWIDAE_LIST_SHEET_NAME in result.sheetnames
        assert QUANTITY_SHEET_NAME not in result.sheetnames
        assert QUANTITY_SHEET_NAME not in result.sheetnames
        assert PUMSAM_SHEET_NAME in result.sheetnames
        assert WAGES_SHEET_NAME in result.sheetnames
        pumsam = result[PUMSAM_SHEET_NAME]
        assert pumsam["A1"].value == "키워드"
        assert pumsam["B1"].value == "명칭"
        assert pumsam["C1"].value == "규격"
        assert pumsam["H1"].value == "품셈근거"
        assert all(
            "공량산출" not in str(pumsam.cell(1, col).value or "")
            for col in range(1, 10)
        )
        wages = result[WAGES_SHEET_NAME]
        assert wages.column_dimensions["A"].width == 17
        assert wages.column_dimensions["B"].width == 17
        assert wages.column_dimensions["C"].width == 40
        assert wages["B4"].number_format == "#,##0"
        assert wages.row_dimensions[4].height == 25

        compare = result[COMPARE_SHEET_NAME]
        assert compare["A1"].value == "단 가 대 비 표"
        assert "품" in str(compare["A3"].value or "").replace(" ", "")
        assert all(
            "코드" not in str(compare.cell(3, col).value or "").replace(" ", "")
            for col in range(1, 22)
        )
        assert compare["A5"].value == "경질비닐전선관_지중"
        assert compare["D5"].value == 200
        assert compare["L5"].value == 200
        assert compare.column_dimensions["E"].width == 5
        assert compare.column_dimensions["K"].width == 5
        assert compare["A6"].value == "경질비닐전선관_노출"
        name_merges = [
            str(range_)
            for range_ in compare.merged_cells.ranges
            if range_.min_col == 1 and range_.min_row >= 5
        ]
        assert name_merges == []

        ilwidae = result[ILWIDAE_SHEET_NAME]
        names = [ilwidae.cell(r, 1).value for r in range(1, 40)]
        assert "전선관부속품비" in names
        assert "잡재료비" in names
        assert "공구손료" in names
        assert "내선전공" in names
        assert "보통인부" in names
        assert ilwidae["A1"].value == "일 위 대 가"
        assert ilwidae["A5"].alignment.horizontal == "left"
        assert ilwidae["E6"].value == "='단가대비표'!L5"
        assert ilwidae["G6"].value == "=0"
        assert "TRUNC" in str(ilwidae["H6"].value)
        assert "G6" in str(ilwidae["H6"].value)
        labor_row = next(r for r in range(5, 30) if ilwidae.cell(r, 1).value == "내선전공")
        fitting_row = next(r for r in range(5, 20) if ilwidae.cell(r, 1).value == "전선관부속품비")
        sundry_row = next(r for r in range(5, 20) if ilwidae.cell(r, 1).value == "잡재료비")
        tool_row = next(r for r in range(5, 25) if ilwidae.cell(r, 1).value == "공구손료")
        assert ilwidae.cell(fitting_row, 2).value == "전선관의 15%"
        assert ilwidae.cell(fitting_row, 6).value == "=TRUNC(F6*0.15,1)"
        assert ilwidae.cell(sundry_row, 2).value == "배관의 2%"
        assert ilwidae.cell(sundry_row, 6).value == "=TRUNC(F6*0.02,1)"
        assert "H" in str(ilwidae.cell(tool_row, 10).value)
        assert "0.03" in str(ilwidae.cell(tool_row, 10).value)
        assert ilwidae.cell(labor_row, 4).value in (None, "")
        assert ilwidae.cell(labor_row, 5).value == "=0"
        assert f"E{labor_row}" in str(ilwidae.cell(labor_row, 6).value)
        assert "VLOOKUP" in str(ilwidae.cell(labor_row, 7).value)
        assert f"G{labor_row}" in str(ilwidae.cell(labor_row, 8).value)
        assert f"D{labor_row}" in str(ilwidae.cell(labor_row, 8).value)
        assert labor_row == 9
        sum_rows = [
            r
            for r in range(5, 25)
            if "합계" in str(ilwidae.cell(r, 1).value or "").replace(" ", "")
        ]
        assert sum_rows
        for price_col in (5, 7, 9, 11):
            assert ilwidae.cell(sum_rows[0], price_col).value in (None, "")
        assert "SUM" in str(ilwidae.cell(sum_rows[0], 6).value)
        assert "SUM" in str(ilwidae.cell(sum_rows[0], 8).value)
        assert "SUM" in str(ilwidae.cell(sum_rows[0], 10).value)
        assert "TRUNC" in str(ilwidae.cell(sum_rows[0], 12).value)
        assert ilwidae["E6"].number_format == "#,##0.00"
        assert ilwidae["F6"].number_format == "#,##0.0"
        assert any("'노임단가'" in str(ilwidae.cell(r, 7).value or "") for r in range(5, 20))
        assert ilwidae.row_dimensions[1].height == 30

        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        assert estimate["A1"].value == "[일위대가목록]"
        assert estimate["A3"].value == "명칭"
        assert estimate["A5"].value == "경질비닐전선관_지중"
        assert estimate["D5"].value == 100
        assert "일위대가" in str(estimate["E5"].value)
        assert f"F{sum_rows[0]}" in str(estimate["E5"].value)
        assert "E5" in str(estimate["F5"].value)
        assert "TRUNC" in str(estimate["F5"].value)
        assert "일위대가" in str(estimate["G5"].value)
        assert f"H{sum_rows[0]}" in str(estimate["G5"].value)
        assert "G5" in str(estimate["H5"].value)
        assert "TRUNC" in str(estimate["H5"].value)
        assert estimate["K5"].value == "=TRUNC(E5+G5+I5,2)"
        assert estimate["L5"].value == "=TRUNC(F5+H5+J5,1)"
        assert estimate["L5"].number_format == "#,##0.0"
        assert estimate.row_dimensions[5].height == 30
        estimate_names = [estimate.cell(r, 1).value for r in range(1, 40)]
        assert not any("부속" in str(value or "") for value in estimate_names)
        assert any(str(value or "").replace(" ", "") == "노무비" for value in estimate_names)
        assert any(str(value or "").replace(" ", "") == "(합계)" for value in estimate_names)
        titles = [
            ilwidae.cell(r, 1).value
            for r in range(1, 40)
            if "호표" in str(ilwidae.cell(r, 1).value or "")
        ]
        assert titles
        assert all("전기" not in str(title) for title in titles)
        remarks = [ilwidae.cell(r, 13).value for r in range(5, 40)]
        assert any(str(value or "").replace(" ", "") == "전기5-1" for value in remarks)
        assert all("품셈" not in str(value or "") for value in remarks)
        assert ilwidae["A5"].font.name == "굴림"
        assert ilwidae["A5"].font.bold is not True
    finally:
        result.close()

    db = tmp_path / "out" / "데이터베이스"
    assert (db / "전기_표준품셈.xlsx").exists()
    assert (db / "통신_표준품셈.xlsx").exists()
    assert (db / "노임단가.xlsx").exists()


def test_match_pumsam_keeps_two_labors() -> None:
    item = LineItem(excel_row=5, name="배선용단자함", spec="10 P 이하", unit="대", qty=1, material_price=1)
    matched = match_pumsam(item, default_pumsam_rows())
    jobs = [row.get("노무명칭") for row in matched]
    assert "내선전공" in jobs
    assert "보통인부" in jobs
    assert len(matched) == 2


def test_write_ilwidae_puts_both_labors_and_keeps_remark_as_ref() -> None:
    from app.ilwidae import write_ilwidae_sheet

    workbook = Workbook()
    sheet = workbook.active
    item = LineItem(excel_row=5, name="배선용단자함", spec="10 P 이하", unit="대", qty=1, material_price=1)
    blocks = write_ilwidae_sheet(sheet, [item], default_pumsam_rows(), default_wage_rows())
    jobs = [sheet.cell(row, 1).value for row in blocks[0].labor_rows]
    assert "내선전공" in jobs
    assert "보통인부" in jobs
    notes = [sheet.cell(row, 13).value for row in blocks[0].labor_rows]
    assert all(note in (None, "") for note in notes)
    material_note = str(sheet.cell(blocks[0].material_row, 13).value or "")
    assert material_note.replace(" ", "").startswith("전기")
    assert "품셈" not in material_note
    kinds = [str(sheet.cell(row, 2).value or "") for row in blocks[0].labor_rows]
    assert all(kind == "일반공사 직종" for kind in kinds)
    assert all(sheet.cell(row, 4).value in (None, "") for row in blocks[0].labor_rows)
    assert all("VLOOKUP" in str(sheet.cell(row, 7).value or "") for row in blocks[0].labor_rows)
    assert all(f"D{row}" in str(sheet.cell(row, 8).value or "") for row in blocks[0].labor_rows)
    workbook.close()


def test_ilwidae_percent_rows_follow_conduit_and_wire_aliases() -> None:
    from app.ilwidae import write_ilwidae_sheet

    workbook = Workbook()
    sheet = workbook.active
    items = [
        LineItem(excel_row=5, name="HI관", spec="16 mm", unit="M", qty=10, material_price=100),
        LineItem(excel_row=6, name="CV케이블", spec="6 ㎟ 이하", unit="M", qty=10, material_price=200),
        LineItem(excel_row=7, name="배선용단자함", spec="10 P 이하", unit="대", qty=1, material_price=1),
    ]
    blocks = write_ilwidae_sheet(sheet, items, default_pumsam_rows(), default_wage_rows())

    def rows_of(block) -> list[tuple[Any, Any, Any]]:
        start = block.material_row
        end = block.sum_row
        return [(sheet.cell(row, 1).value, sheet.cell(row, 2).value, sheet.cell(row, 6).value) for row in range(start, end)]

    conduit_rows = rows_of(blocks[0])
    assert ("전선관부속품비", "전선관의 15%", f"=TRUNC(F{blocks[0].material_row}*0.15,1)") in conduit_rows
    assert ("잡재료비", "배관의 2%", f"=TRUNC(F{blocks[0].material_row}*0.02,1)") in conduit_rows
    assert sheet.cell(blocks[0].sum_row - 1, 1).value == "공구손료"

    wire_rows = rows_of(blocks[1])
    assert all(name != "전선관부속품비" for name, _spec, _amount in wire_rows)
    assert ("잡재료비", "배선의 2%", f"=TRUNC(F{blocks[1].material_row}*0.02,1)") in wire_rows
    wire_jobs = [sheet.cell(row, 1).value for row in blocks[1].labor_rows]
    assert "내선전공" not in wire_jobs
    assert "케이블전공" in wire_jobs
    assert "보통인부" in wire_jobs
    labor = next(row for row in blocks[1].labor_rows if sheet.cell(row, 1).value == "케이블전공")
    assert "케이블전공" in str(sheet.cell(labor, 7).value)
    assert "노임단가" in str(sheet.cell(labor, 7).value)

    box_rows = rows_of(blocks[2])
    assert all(name not in {"전선관부속품비", "잡재료비"} for name, _spec, _amount in box_rows)
    assert "내선전공" in [sheet.cell(row, 1).value for row in blocks[2].labor_rows]
    assert "보통인부" in [sheet.cell(row, 1).value for row in blocks[2].labor_rows]
    assert sheet.cell(blocks[2].sum_row - 1, 1).value == "공구손료"
    tool = blocks[2].sum_row - 1
    tool_formula = str(sheet.cell(tool, 10).value)
    assert "0.03" in tool_formula
    for labor_row in blocks[2].labor_rows:
        assert f"H{labor_row}" in tool_formula
    workbook.close()


def test_unmatched_pumsam_does_not_use_similar_names() -> None:
    item = LineItem(excel_row=5, name="특수커넥터함", spec="일반", unit="개", qty=1, material_price=1)
    assert match_pumsam(item, default_pumsam_rows()) == []


def test_bare_conduit_uses_maip_not_exposed() -> None:
    rows = default_pumsam_rows()
    matched = match_pumsam(
        LineItem(excel_row=5, name="경질비닐전선관", spec="HI 28 mm", unit="M", qty=1, material_price=1),
        rows,
    )
    assert len(matched) == 1
    assert matched[0]["노무명칭"] == "내선전공"
    assert matched[0]["품셈"] == 0.08
    assert matched[0]["할증%"] == 100
    exposed = match_pumsam(
        LineItem(excel_row=6, name="경질비닐전선관_노출", spec="HI 28 mm", unit="M", qty=1, material_price=1),
        rows,
    )
    assert len(exposed) == 1
    assert exposed[0]["할증%"] == 120
    assert "*1.2" in labor_qty_formula(exposed[0])
    buried = match_pumsam(
        LineItem(excel_row=7, name="경질비닐전선관_지중", spec="HI 16 mm", unit="M", qty=1, material_price=1),
        rows,
    )
    assert len(buried) == 1
    assert buried[0]["할증%"] == 70
    assert buried[0]["품셈"] == 0.05


def test_unmatched_item_gets_one_fallback_labor(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "단가대비표"
    sheet["A1"] = "단 가 대 비 표"
    sheet["A3"] = "품명"
    sheet["B3"] = "규격"
    sheet["C3"] = "단위"
    sheet["D3"] = "수량"
    sheet["A5"] = "특수커넥터함"
    sheet["B5"] = "일반"
    sheet["C5"] = "개"
    sheet["L5"] = 1021
    source = tmp_path / "단가대비표.xlsx"
    workbook.save(source)
    workbook.close()
    dest = save_result_workbook(unit_price_path=source, dest_dir=tmp_path / "out")
    result = load_workbook(dest, data_only=False)
    try:
        ilwidae = result[ILWIDAE_SHEET_NAME]
        title = next(r for r in range(5, 20) if "호표" in str(ilwidae.cell(r, 1).value or ""))
        total = next(
            r
            for r in range(title, 30)
            if "합계" in str(ilwidae.cell(r, 1).value or "").replace(" ", "")
        )
        jobs = [ilwidae.cell(r, 1).value for r in range(title + 1, total)]
        assert jobs.count("내선전공") == 1
        assert "보통인부" not in jobs
        labor_row = title + 2
        assert ilwidae.cell(labor_row, 1).value == "내선전공"
        assert ilwidae.cell(labor_row, 4).value in (None, "")
        assert "VLOOKUP" in str(ilwidae.cell(labor_row, 7).value)
        assert f"D{labor_row}" in str(ilwidae.cell(labor_row, 8).value)
    finally:
        result.close()


def test_unmatched_telecom_item_gets_one_telecom_labor(tmp_path: Path) -> None:
    from app.discipline import TELECOM

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "단가대비표"
    sheet["A1"] = "단 가 대 비 표"
    sheet["A3"] = "품명"
    sheet["B3"] = "규격"
    sheet["C3"] = "단위"
    sheet["D3"] = "수량"
    sheet["A5"] = "강제전선관"
    sheet["B5"] = "아연도 16 mm"
    sheet["C5"] = "M"
    sheet["L5"] = 1338
    source = tmp_path / "단가대비표.xlsx"
    workbook.save(source)
    workbook.close()
    dest = save_result_workbook(
        unit_price_path=source,
        dest_dir=tmp_path / "out",
        discipline=TELECOM,
    )
    result = load_workbook(dest, data_only=False)
    try:
        ilwidae = result[ILWIDAE_SHEET_NAME]
        title = next(r for r in range(5, 20) if "호표" in str(ilwidae.cell(r, 1).value or ""))
        total = next(
            r
            for r in range(title, 30)
            if "합계" in str(ilwidae.cell(r, 1).value or "").replace(" ", "")
        )
        jobs = [ilwidae.cell(r, 1).value for r in range(title + 1, total)]
        assert jobs.count("통신내선공") == 1
        assert "내선전공" not in jobs
        assert "보통인부" not in jobs
    finally:
        result.close()


def test_default_wages_include_trades() -> None:
    rows = {row["직종"]: row for row in default_wage_rows()}
    assert rows["내선전공"]["노임단가"] == 276108
    assert rows["보통인부"]["노임단가"] == 172698
    assert rows["저압케이블전공"]["노임단가"] == 306274
    assert "2026" in str(rows["내선전공"]["비고"])


def test_sample_unit_price_skips_header_and_empty_qty(tmp_path: Path) -> None:
    sample = Path("/home/ubuntu/.cursor/projects/workspace/uploads/_________797b.xlsx")
    if not sample.exists():
        return
    dest = save_result_workbook(unit_price_path=sample, dest_dir=tmp_path / "out")
    result = load_workbook(dest, data_only=False)
    try:
        ilwidae = result[ILWIDAE_SHEET_NAME]
        titles = [
            str(ilwidae.cell(r, 1).value)
            for r in range(1, 30)
            if "호표" in str(ilwidae.cell(r, 1).value or "")
        ]
        assert titles
        assert all("품 명" not in title and "품명" not in title.replace(" ", "") for title in titles)
        assert "강제전선관" in titles[0]
        assert all("전기" not in title for title in titles)
        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        compare = result[COMPARE_SHEET_NAME]
        assert all(
            "코드" not in str(compare.cell(3, col).value or "").replace(" ", "")
            for col in range(1, 22)
        )
        assert compare["A5"].value == "강제전선관"
        assert compare["A6"].value == "강제전선관"
        name_merges = [
            str(range_)
            for range_ in compare.merged_cells.ranges
            if range_.min_col == 1 and range_.min_row >= 5
        ]
        assert name_merges == []
        assert compare["E5"].value == 1202
        assert estimate["A5"].value == "강제전선관"
        assert estimate["D5"].value in (None, "")
        assert estimate["C5"].value == "M"
        assert "일위대가" in str(estimate["E5"].value or "")
        assert "E5" in str(estimate["F5"].value or "")
    finally:
        result.close()


def test_reverse_estimate_builds_unit_price(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일위대가목록"
    sheet["A1"] = "[일위대가목록]"
    sheet["A2"] = "명칭"
    sheet["B2"] = "규격"
    sheet["C2"] = "단위"
    sheet["D2"] = "수량"
    sheet["E2"] = "재료비"
    sheet["E3"] = "단가"
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet.merge_cells("C2:C3")
    sheet.merge_cells("D2:D3")
    sheet.merge_cells("E2:F2")
    sheet["A4"] = "1. 전기공사"
    sheet["A5"] = "강제전선관"
    sheet["B5"] = "아연도 16 mm"
    sheet["C5"] = "M"
    sheet["D5"] = 80
    sheet["E5"] = 3374
    source = tmp_path / "내역서.xlsx"
    workbook.save(source)
    workbook.close()

    dest = save_result_workbook(estimate_path=source, dest_dir=tmp_path / "out")
    result = load_workbook(dest, data_only=False)
    try:
        assert dest.name.startswith("단가대비표_결과_")
        assert result.sheetnames[0] == COMPARE_SHEET_NAME
        assert ILWIDAE_LIST_SHEET_NAME in result.sheetnames
        assert QUANTITY_SHEET_NAME not in result.sheetnames
        compare = result[COMPARE_SHEET_NAME]
        assert compare["A1"].value == "단 가 대 비 표"
        assert "품" in str(compare["A3"].value or "").replace(" ", "")
        assert all(
            "코드" not in str(compare.cell(3, col).value or "").replace(" ", "")
            for col in range(1, 22)
        )
        assert compare["A5"].value == "강제전선관"
        assert compare["D5"].value == 3374
        assert compare["L5"].value == 3374
        assert compare.row_dimensions[5].height == 30
    finally:
        result.close()


def test_telecom_pumsam_does_not_pull_electric_labor(tmp_path: Path) -> None:
    from app.discipline import TELECOM
    from app.pumsam import load_pumsam_database, pumsam_filename

    electric = {row.get("노무명칭") for row in default_pumsam_rows("전기")}
    telecom = {row.get("노무명칭") for row in default_pumsam_rows(TELECOM)}
    assert "내선전공" in electric
    assert "내선전공" not in telecom
    assert "통신내선공" in telecom
    assert "통신케이블공" in telecom

    source = tmp_path / "단가대비표.xlsx"
    _write_compare(source)
    dest = save_result_workbook(
        unit_price_path=source,
        dest_dir=tmp_path / "out",
        discipline=TELECOM,
    )
    result = load_workbook(dest, data_only=False)
    try:
        ilwidae = result[ILWIDAE_SHEET_NAME]
        names = [ilwidae.cell(r, 1).value for r in range(1, 40)]
        assert "통신내선공" in names
        assert "내선전공" not in names
        remarks = [str(ilwidae.cell(r, 13).value or "").replace(" ", "") for r in range(5, 40)]
        assert any(value in {"전기5-1", "통신3-1-1"} for value in remarks)
        assert all("품셈" not in str(ilwidae.cell(r, 13).value or "") for r in range(5, 40))
        assert any(str(ilwidae.cell(r, 2).value or "") == "일반공사 직종" for r in range(5, 40))
        assert all("할증" not in str(ilwidae.cell(r, 2).value or "") for r in range(5, 40))
        assert QUANTITY_SHEET_NAME not in result.sheetnames
    finally:
        result.close()

    db = tmp_path / "out" / "데이터베이스"
    assert (db / pumsam_filename(TELECOM)).exists()
    assert (db / "전기_표준품셈.xlsx").exists()
    loaded = load_pumsam_database(tmp_path / "out", TELECOM)
    jobs = {row.get("노무명칭") for row in loaded if "경질비닐전선관" in str(row.get("명칭") or "")}
    assert "통신내선공" in jobs
    assert "내선전공" not in jobs


def test_quantity_mode_keeps_estimate_parts(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일위대가목록"
    sheet["A1"] = "[일위대가목록]"
    sheet["A3"] = "명칭"
    sheet["B3"] = "규격"
    sheet["C3"] = "단위"
    sheet["D3"] = "수량"
    sheet["E3"] = "재료비"
    sheet["E4"] = "단가"
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:C4")
    sheet.merge_cells("D3:D4")
    sheet.merge_cells("E3:F3")
    sheet["A5"] = "1. 전기공사"
    sheet["A6"] = "강제전선관"
    sheet["B6"] = "아연도 16 mm"
    sheet["C6"] = "M"
    sheet["A7"] = "1. 전열설비공사"
    sheet["A8"] = "경질비닐전선관"
    sheet["B8"] = "HI 16 mm"
    sheet["C8"] = "M"
    source = tmp_path / "내역서_파트.xlsx"
    workbook.save(source)
    workbook.close()

    dest = save_result_workbook(estimate_path=source, dest_dir=tmp_path / "out", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        assert dest.name.startswith("공량산출_결과_")
        assert COMPARE_SHEET_NAME not in result.sheetnames
        assert QUANTITY_SHEET_NAME in result.sheetnames
        qty = result[QUANTITY_SHEET_NAME]
        names = [qty.cell(r, 2).value for r in range(4, 10)]
        assert "1. 전기공사" in names
        assert "1. 전열설비공사" in names
        assert names.index("1. 전기공사") < names.index("강제전선관") < names.index("1. 전열설비공사")
        assert qty["B8"].value == "경질비닐전선관"
        assert qty["H8"].value in (None, "")
        assert qty["I8"].value in (None, "")
        assert qty["J8"].value == 100
        assert "G8*I8*(J8/100)" in str(qty["K8"].value)
        assert "VLOOKUP" not in str(qty["H8"].value or "")
        assert "SUMPRODUCT" not in str(qty["K8"].value or "")
        assert qty["A3"].value == "품목"
        from app.pumsam import PUMSAM_COLUMN_WIDTHS

        pumsam = result[PUMSAM_SHEET_NAME]
        assert pumsam["A1"].value == "키워드"
        assert pumsam.column_dimensions["A"].width == PUMSAM_COLUMN_WIDTHS[0]
        assert pumsam.column_dimensions["C"].width == PUMSAM_COLUMN_WIDTHS[2]
        wages = result[WAGES_SHEET_NAME]
        assert wages.column_dimensions["A"].width == 17
        assert wages.column_dimensions["B"].width == 17
        assert wages.column_dimensions["C"].width == 40
        assert wages["B4"].number_format == "#,##0"
        assert wages.row_dimensions[4].height == 25
    finally:
        result.close()


def test_compare_keeps_page_and_pps_prices(tmp_path: Path) -> None:
    sample = Path("/home/ubuntu/.cursor/projects/workspace/uploads/_________8473.xlsx")
    if not sample.exists():
        return
    dest = save_result_workbook(unit_price_path=sample, dest_dir=tmp_path / "out")
    result = load_workbook(dest, data_only=False)
    try:
        compare = result[COMPARE_SHEET_NAME]
        assert compare["A5"].value == "강제전선관"
        assert compare["D5"].value == 3374
        assert compare["E5"].value == 1202
        assert compare.column_dimensions["E"].width == 5
        assert compare.column_dimensions["G"].width == 5
        assert compare.column_dimensions["I"].width == 5
        assert compare.column_dimensions["K"].width == 5
        assert compare["L5"].value == 3374
        rows = {
            (compare.cell(r, 1).value, str(compare.cell(r, 2).value or "")): r
            for r in range(5, 40)
        }
        d30 = rows.get(("관로구방수장치", "D30"))
        assert d30
        assert compare.cell(d30, 4).value in (None, "")
        assert compare.cell(d30, 6).value == 14350
        assert compare.cell(d30, 12).value == 14350
        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        assert "일위대가" in str(estimate["E5"].value)
        assert "E5" in str(estimate["F5"].value)
    finally:
        result.close()


def _merged_overlap(left, right) -> bool:
    return not (
        left.max_row < right.min_row
        or left.min_row > right.max_row
        or left.max_col < right.min_col
        or left.min_col > right.max_col
    )


def test_quantity_keeps_sheet_with_overlapping_source_merges(tmp_path: Path) -> None:
    from app.excel_io import _without_overlapping_merges
    from app.estimate_parse import EstimateSheet
    from app.pipeline import build_result_workbook
    from app.pumsam import default_pumsam_rows
    from app.wages import default_wage_rows

    kept = _without_overlapping_merges([(1, 1, 1, 13), (1, 1, 2, 1), (2, 1, 3, 1)])
    assert kept[0] == (1, 1, 1, 13)
    assert (1, 1, 2, 1) not in kept

    filled = [
        ["[내역서 ]"] + [None] * 12,
        ["명칭", "규격", "단위", "수량", "재료비"] + [None] * 8,
        [None, None, None, None, "단가", "금액"] + [None] * 7,
        ["1. 옥외전기공사"] + [None] * 12,
        ["경질비닐전선관_지중", "HI 16 mm", "M", 10, 200, "=D5*E5"] + [None] * 7,
    ]
    estimate = EstimateSheet(
        filled=filled,
        raw=filled,
        merges=[(1, 1, 1, 13), (1, 1, 2, 1), (2, 1, 3, 1), (2, 5, 2, 6)],
    )
    workbook = build_result_workbook(
        compare=None,
        estimate=estimate,
        pumsam_rows=default_pumsam_rows(),
        wage_rows=default_wage_rows(),
        mode="quantity",
    )
    dest = tmp_path / "out.xlsx"
    workbook.save(dest)
    workbook.close()
    result = load_workbook(dest, data_only=False)
    try:
        assert QUANTITY_SHEET_NAME in result.sheetnames
        assert ILWIDAE_LIST_SHEET_NAME in result.sheetnames
        qty = result[QUANTITY_SHEET_NAME]
        assert qty["B5"].value == "경질비닐전선관_지중"
        for sheet in result.worksheets:
            ranges = list(sheet.merged_cells.ranges)
            for i, left in enumerate(ranges):
                for right in ranges[i + 1 :]:
                    assert not _merged_overlap(left, right), f"{sheet.title}: {left} vs {right}"
    finally:
        result.close()


def test_quantity_strips_external_workbook_formulas(tmp_path: Path) -> None:
    from app.estimate_parse import list_sheet_titles

    sample = Path("/home/ubuntu/.cursor/projects/workspace/uploads/_______d25a.xlsx")
    if not sample.exists():
        return
    titles = list_sheet_titles(sample)
    if ILWIDAE_LIST_SHEET_NAME not in titles and ESTIMATE_SHEET_NAME not in titles:
        return
    if ILWIDAE_LIST_SHEET_NAME not in titles:
        return
    dest = save_result_workbook(estimate_path=sample, dest_dir=tmp_path / "out", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        assert QUANTITY_SHEET_NAME in result.sheetnames
        for row in estimate.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    assert "[" not in value
                    assert "노임산출서" not in value
        assert estimate["F5"].value == "=D5*E5"
        assert estimate["D214"].value in (0, None)
    finally:
        result.close()


def test_quantity_skips_sundry_form_formulas(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일위대가목록"
    sheet["A1"] = "[일위대가목록]"
    sheet["A3"] = "명칭"
    sheet["B3"] = "규격"
    sheet["C3"] = "단위"
    sheet["D3"] = "수량"
    sheet["E3"] = "재료비"
    sheet["E4"] = "단가"
    sheet["A5"] = "1. 전기공사"
    sheet["A6"] = "강제전선관"
    sheet["B6"] = "아연도 16 mm"
    sheet["C6"] = "M"
    sheet["D6"] = 10
    sheet["A8"] = "[ 배관 부속재 ]"
    sheet["B8"] = "전선관의 15 %"
    sheet["C8"] = "식"
    sheet["D8"] = 1
    sheet["A9"] = "노 무 비"
    sheet["B9"] = "내선전공"
    sheet["C9"] = "인"
    sheet["A10"] = "( 합 계 )"
    source = tmp_path / "내역서_양식.xlsx"
    workbook.save(source)
    workbook.close()

    dest = save_result_workbook(estimate_path=source, dest_dir=tmp_path / "out", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        qty = result[QUANTITY_SHEET_NAME]
        assert qty["B6"].value == "강제전선관"
        assert qty["H6"].value in (None, "")
        assert qty["I6"].value in (None, "")
        assert "G6*I6*(J6/100)" in str(qty["K6"].value)
        assert qty["B8"].value == "[ 배관 부속재 ]"
        assert qty["C8"].value == "전선관의 15 %"
        assert qty["H8"].value in (None, "")
        assert qty["K8"].value in (None, "")
        assert qty["B9"].value == "노 무 비"
        assert qty["H9"].value in (None, "")
        assert qty["B10"].value == "( 합 계 )"
        assert qty["H10"].value in (None, "")
    finally:
        result.close()


def test_quantity_keeps_compare_and_ilwidae_from_forward_result(tmp_path: Path) -> None:
    source = tmp_path / "단가대비표.xlsx"
    _write_compare(source)
    forward = save_result_workbook(unit_price_path=source, dest_dir=tmp_path / "fwd")
    dest = save_result_workbook(estimate_path=forward, dest_dir=tmp_path / "qty", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        assert dest.name.startswith("공량산출_결과_")
        assert COMPARE_SHEET_NAME in result.sheetnames
        assert ILWIDAE_SHEET_NAME in result.sheetnames
        assert ILWIDAE_LIST_SHEET_NAME in result.sheetnames
        assert QUANTITY_SHEET_NAME in result.sheetnames
        qty = result[QUANTITY_SHEET_NAME]
        assert qty["B5"].value == "경질비닐전선관_지중"
        assert "일위대가목록" in str(qty["G5"].value)
        compare = result[COMPARE_SHEET_NAME]
        assert compare["A5"].value == "경질비닐전선관_지중"
        ilwidae = result[ILWIDAE_SHEET_NAME]
        titles = [
            str(ilwidae.cell(r, 1).value)
            for r in range(1, 40)
            if "호표" in str(ilwidae.cell(r, 1).value or "")
        ]
        assert titles
        assert "( 호표 1 )" in titles[0]
        assert "( 호표 2 )" in titles[1]
        from app.ilwidae import parse_ilwidae_blocks

        blocks = parse_ilwidae_blocks(ilwidae)
        first = next(block for block in blocks if str(block.item.name) == "경질비닐전선관_지중")
        assert first.labor_rows
        assert qty["H5"].value == f"='일위대가'!A{first.labor_rows[0]}"
        assert qty["I5"].value == "=" + "+".join(f"'일위대가'!D{r}" for r in first.labor_rows)
        assert qty["J5"].value == 70
        assert qty["L5"].value == f"='일위대가'!M{first.material_row}"
        assert "G5*I5*(J5/100)" in str(qty["K5"].value)
        assert all(ilwidae.cell(row, 4).value in (None, "") for row in first.labor_rows)
        box = next(block for block in blocks if "단자함" in str(block.item.name or ""))
        qty_row = next(r for r in range(5, 20) if qty.cell(r, 2).value == "배선용단자함")
        assert len(box.labor_rows) == 2
        assert qty.cell(qty_row, 9).value == "=" + "+".join(
            f"'일위대가'!D{r}" for r in box.labor_rows
        )
    finally:
        result.close()


def test_quantity_requires_ilwidae_list_sheet(tmp_path: Path) -> None:
    import pytest

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "내역서"
    sheet["A1"] = "[내역서]"
    sheet["A3"] = "명칭"
    sheet["A5"] = "강제전선관"
    source = tmp_path / "내역서만.xlsx"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="일위대가목록 시트"):
        save_result_workbook(estimate_path=source, dest_dir=tmp_path / "out", mode="quantity")


def test_quantity_renumbers_ho_and_reads_pumsam_from_ilwidae(tmp_path: Path) -> None:
    workbook = Workbook()
    listing = workbook.active
    listing.title = "일위대가목록"
    listing["A1"] = "[일위대가목록]"
    listing["A3"] = "명칭"
    listing["B3"] = "규격"
    listing["C3"] = "단위"
    listing["D3"] = "수량"
    listing["A5"] = "강제전선관"
    listing["B5"] = "아연도 16 mm"
    listing["C5"] = "M"
    listing["A6"] = "경질비닐전선관"
    listing["B6"] = "HI 16 mm"
    listing["C6"] = "M"

    ilwidae = workbook.create_sheet("일위대가")
    ilwidae["A1"] = "일 위 대 가"
    ilwidae["A5"] = "경질비닐전선관 HI 16 mm  ( 호표 9 )"
    ilwidae["A6"] = "경질비닐전선관"
    ilwidae["B6"] = "HI 16 mm"
    ilwidae["C6"] = "M"
    ilwidae["A7"] = "내선전공"
    ilwidae["B7"] = "일반공사 직종"
    ilwidae["C7"] = "인"
    ilwidae["D7"] = 0.05
    ilwidae["M6"] = "전기 5-1"
    ilwidae["A8"] = " [ 합          계 ]"
    ilwidae["A10"] = "강제전선관 아연도 16 mm  ( 호표 3 )"
    ilwidae["A11"] = "강제전선관"
    ilwidae["B11"] = "아연도 16 mm"
    ilwidae["C11"] = "M"
    ilwidae["A12"] = "내선전공"
    ilwidae["C12"] = "인"
    ilwidae["D12"] = 0.08
    ilwidae["M11"] = "전기 5-1"
    ilwidae["A13"] = " [ 합          계 ]"
    source = tmp_path / "수동수정.xlsx"
    workbook.save(source)
    workbook.close()

    dest = save_result_workbook(estimate_path=source, dest_dir=tmp_path / "out", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        ilwidae = result[ILWIDAE_SHEET_NAME]
        titles = [
            str(ilwidae.cell(r, 1).value)
            for r in range(1, 20)
            if "호표" in str(ilwidae.cell(r, 1).value or "")
        ]
        assert titles == [
            "경질비닐전선관 HI 16 mm  ( 호표 1 )",
            "강제전선관 아연도 16 mm  ( 호표 2 )",
        ]
        qty = result[QUANTITY_SHEET_NAME]
        assert qty["B5"].value == "강제전선관"
        assert qty["H5"].value == "='일위대가'!A12"
        assert qty["I5"].value == "='일위대가'!D12"
        assert qty["L5"].value == "='일위대가'!M11"
        assert qty["B6"].value == "경질비닐전선관"
        assert qty["I6"].value == "='일위대가'!D7"
        assert "품셈표" not in str(qty["I5"].value)
        assert "VLOOKUP" not in str(qty["I5"].value)
        assert qty["K5"].value == (
            '=IF(OR(G5="",I5="",G5*I5=0),"",G5*I5*(J5/100))'
        )
    finally:
        result.close()

