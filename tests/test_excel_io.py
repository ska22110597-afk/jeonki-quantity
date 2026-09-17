from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.excel_io import (
    FORM_ROW_HEIGHT,
    ILWIDAE_LIST_SHEET_NAME,
    NUMBER_FORMAT,
    QUANTITY_SHEET_NAME,
    ROW_HEIGHT,
    concat_formula,
    decided_qty_formula,
    read_unit_price_table,
    save_result_workbook,
    source_qty_formula,
)
from app.estimate_parse import first_data_row_number, is_section_row, is_sundry_form_row, load_estimate_sheet, lookup_key
from app.pumsam import PUMSAM_SHEET_NAME, import_pumsam_file


def _write_estimate(path: Path, *, with_merge: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일위대가목록"
    sheet["A1"] = "[일위대가목록]"
    sheet["A2"] = "명칭"
    sheet["B2"] = "규격"
    sheet["C2"] = "단위"
    sheet["D2"] = "수량"
    sheet["E2"] = "재료비"
    sheet["G2"] = "노무비"
    sheet["E3"] = "단가"
    sheet["F3"] = "금액"
    sheet["G3"] = "단가"
    sheet["H3"] = "금액"
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet.merge_cells("C2:C3")
    sheet.merge_cells("D2:D3")
    sheet.merge_cells("E2:F2")
    sheet.merge_cells("G2:H2")
    sheet["A4"] = "1. 옥외전기공사"
    sheet["A5"] = "경질비닐전선관_지중"
    sheet["B5"] = "HI 16 mm"
    sheet["C5"] = "M"
    sheet["D5"] = 10
    sheet["E5"] = 200
    sheet["F5"] = "=D5*E5"
    sheet["A6"] = None if with_merge else "경질비닐전선관_지중"
    sheet["B6"] = "HI 22 mm"
    sheet["C6"] = "M"
    sheet["D6"] = 8
    sheet["A7"] = "경질비닐전선관_노출"
    sheet["B7"] = "HI 16 mm"
    sheet["C7"] = "M"
    sheet["D7"] = 4
    if with_merge:
        sheet.merge_cells("A5:A6")
    workbook.save(path)
    workbook.close()


def _fill_rgb(cell) -> str:
    color = cell.fill.fgColor
    rgb = getattr(color, "rgb", None)
    return str(rgb).upper() if rgb is not None else ""


def test_section_row_detection() -> None:
    assert is_section_row("1. 옥외전기공사", None, None)
    assert not is_section_row("경질비닐전선관_지중", "HI 16 mm", "M")
    assert is_sundry_form_row("[ 배관 부속재 ]")
    assert is_sundry_form_row("노 무 비")
    assert is_sundry_form_row("( 합 계 )")
    assert not is_sundry_form_row("강제전선관")


def test_read_only_skips_title_keeps_items(tmp_path: Path) -> None:
    source = tmp_path / "내역서.xlsx"
    _write_estimate(source)
    before = source.read_bytes()
    rows = read_unit_price_table(source)
    assert source.read_bytes() == before
    assert rows[0][0] == "명칭"
    assert rows[2][0] == "1. 옥외전기공사"
    assert rows[3][0] == "경질비닐전선관_지중"


def test_merged_name_fills_without_rewriting_source(tmp_path: Path) -> None:
    source = tmp_path / "내역서_병합.xlsx"
    _write_estimate(source, with_merge=True)
    before = source.read_bytes()
    rows = read_unit_price_table(source)
    assert source.read_bytes() == before
    assert rows[3][0] == "경질비닐전선관_지중"
    assert rows[4][0] == "경질비닐전선관_지중"


def test_full_grid_keeps_original_row_numbers(tmp_path: Path) -> None:
    source = tmp_path / "내역서.xlsx"
    _write_estimate(source)
    estimate = load_estimate_sheet(source)
    assert estimate.filled[0][0] == "[일위대가목록]"
    assert first_data_row_number(estimate.filled) == 4
    assert estimate.filled[4][0] == "경질비닐전선관_지중"
    assert estimate.filled[4][3] == 10


def test_three_sheets_sample_layout_and_same_row_formulas(tmp_path: Path) -> None:
    source = tmp_path / "내역서.xlsx"
    _write_estimate(source)
    dest_dir = tmp_path / "결과"
    dest = save_result_workbook(source, dest_dir=dest_dir, mode="quantity")

    assert dest.name.startswith("공량산출_결과_")
    assert dest.parent == dest_dir

    result = load_workbook(dest, data_only=False)
    try:
        assert ILWIDAE_LIST_SHEET_NAME in result.sheetnames
        assert PUMSAM_SHEET_NAME in result.sheetnames
        assert QUANTITY_SHEET_NAME in result.sheetnames
        assert result.sheetnames[0] == ILWIDAE_LIST_SHEET_NAME
        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        assert estimate["A1"].value == "[일위대가목록]"
        assert estimate["A2"].value == "명칭"
        assert estimate["D5"].value == 10
        assert estimate["F5"].value == "=D5*E5"
        assert estimate.row_dimensions[5].height == FORM_ROW_HEIGHT
        assert estimate["A5"].alignment.wrap_text is not True
        assert "FFFFFF" in _fill_rgb(estimate["A2"])
        assert str(estimate.sheet_properties.tabColor.rgb).upper().endswith("FFFFFF")

        pumsam = result[PUMSAM_SHEET_NAME]
        assert pumsam["A1"].value == "품 셈 표"
        assert pumsam["A3"].value == "품목"
        assert pumsam["E4"].value == "명칭"
        assert pumsam["B5"].value
        names = [pumsam.cell(r, 2).value for r in range(5, pumsam.max_row + 1)]
        assert "경질비닐전선관_지중" in names
        assert str(pumsam["A5"].value).startswith("=CONCATENATE")
        extra_labor = False
        for row_idx in range(5, pumsam.max_row + 1):
            if pumsam.cell(row_idx, 2).value in (None, "") and pumsam.cell(row_idx, 3).value not in (None, ""):
                extra_labor = True
                assert pumsam.cell(row_idx, 1).value in (None, "")
                assert pumsam.cell(row_idx, 5).value not in (None, "")
                break
        assert extra_labor
        assert pumsam.column_dimensions["A"].width == 50
        assert pumsam.column_dimensions["C"].width == 50
        assert pumsam.row_dimensions[5].height == ROW_HEIGHT
        assert "FFFFFF" in _fill_rgb(pumsam["A3"])

        qty = result[QUANTITY_SHEET_NAME]
        assert qty["A1"].value == "공 량 산 출 서"
        assert qty["B2"].value == "명칭"
        assert qty["E3"].value == "결정수량"
        assert qty["B4"].value == "1. 옥외전기공사"
        assert qty["B5"].value == "경질비닐전선관_지중"
        assert qty["C5"].value == "HI 16 mm"
        assert qty["A5"].value == concat_formula(5)
        assert qty["E5"].value == decided_qty_formula(5)
        assert qty["F5"].value == "=0"
        assert qty["G5"].value == source_qty_formula("D", 5, ILWIDAE_LIST_SHEET_NAME)
        assert qty["G5"].value == "='일위대가목록'!D5"
        assert qty["H5"].value in (None, "")
        assert qty["I5"].value in (None, "")
        assert qty["J5"].value == 100
        assert qty["K5"].value == (
            '=IF(OR(G5="",I5="",G5*I5=0),"",G5*I5*(J5/100))'
        )
        assert "VLOOKUP" not in str(qty["H5"].value or "")
        assert "SUMPRODUCT" not in str(qty["K5"].value or "")
        assert qty["A2"].value == "품목"
        assert "G5" in str(qty["K5"].value)
        assert qty["K5"].number_format == NUMBER_FORMAT
        assert qty.row_dimensions[5].height == ROW_HEIGHT
        assert qty["B5"].alignment.wrap_text is not True
        assert "원본" not in str(qty["A2"].value or "")
        assert "FFFFFF" in _fill_rgb(qty["B2"])
        names = [qty.cell(r, 2).value for r in range(4, 8)]
        assert "1. 옥외전기공사" in names
        assert lookup_key("경질비닐전선관_지중", "HI 16 mm") == "경질비닐전선관_지중HI16mm"
    finally:
        result.close()


def test_merged_source_does_not_error(tmp_path: Path) -> None:
    source = tmp_path / "내역서_병합.xlsx"
    _write_estimate(source, with_merge=True)
    dest = save_result_workbook(source, dest_dir=tmp_path / "out", mode="quantity")
    result = load_workbook(dest, data_only=False)
    try:
        qty = result[QUANTITY_SHEET_NAME]
        assert qty["B5"].value == "경질비닐전선관_지중"
        assert qty["B6"].value == "경질비닐전선관_지중"
        assert qty["G6"].value == "='일위대가목록'!D6"
        estimate = result[ILWIDAE_LIST_SHEET_NAME]
        assert estimate["A5"].value == "경질비닐전선관_지중"
    finally:
        result.close()


def test_custom_dest_directory(tmp_path: Path) -> None:
    source = tmp_path / "내역서.xlsx"
    _write_estimate(source)
    chosen = tmp_path / "내가지정한폴더"
    dest = save_result_workbook(source, dest_dir=chosen)
    assert dest.parent == chosen
    assert dest.exists()


def test_import_two_row_pumsam_sample() -> None:
    sample = Path("/home/ubuntu/.cursor/projects/workspace/uploads/_______bfba.xlsx")
    if not sample.exists():
        return
    rows = import_pumsam_file(sample)
    names = {str(row.get("명칭")) for row in rows}
    assert "경질비닐전선관_지중" in names
    assert "경질비닐전선관_노출" in names
    assert len(rows) >= 20
