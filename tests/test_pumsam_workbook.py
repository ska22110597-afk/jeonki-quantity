from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.estimate_parse import display_keyword
from app.pumsam import (
    PUMSAM_COLUMN_WIDTHS,
    PUMSAM_DISPLAY_HEADERS,
    PUMSAM_ROW_HEIGHT,
    default_pumsam_rows,
    save_pumsam_database,
)


def _fill_rgb(cell) -> str:
    color = cell.fill.fgColor
    rgb = getattr(color, "rgb", None)
    return str(rgb).upper() if rgb is not None else ""


def test_pumsam_workbook_layout_electric_and_telecom(tmp_path: Path) -> None:
    for disc in ("전기", "통신"):
        rows = default_pumsam_rows(disc)[:40]
        path = save_pumsam_database(rows, directory=tmp_path, discipline=disc)
        workbook = load_workbook(path)
        try:
            sheet = workbook["품셈표"]
            assert [sheet.cell(1, col).value for col in range(1, 9)] == list(PUMSAM_DISPLAY_HEADERS)
            assert sheet["A1"].value == "키워드"
            assert sheet.row_dimensions[1].height == PUMSAM_ROW_HEIGHT
            assert sheet.row_dimensions[2].height == PUMSAM_ROW_HEIGHT
            assert sheet["A1"].font.name == "굴림"
            assert sheet["B2"].font.name == "굴림"
            assert sheet["A1"].alignment.horizontal == "center"
            assert sheet["B1"].alignment.horizontal == "center"
            assert "B7DEE8" in _fill_rgb(sheet["A1"])
            assert "FFFFFF" in _fill_rgb(sheet["A2"]) or _fill_rgb(sheet["A2"]).endswith("FFFFFF")
            assert sheet["B2"].alignment.horizontal == "left"
            assert sheet["C2"].alignment.horizontal == "left"
            assert sheet["D2"].alignment.horizontal == "center"
            assert sheet["E2"].alignment.horizontal == "center"
            assert sheet["F2"].alignment.horizontal == "right"
            assert sheet["G2"].alignment.horizontal == "center"
            assert sheet["H2"].alignment.horizontal == "center"
            letters = "ABCDEFGH"
            for letter, width in zip(letters, PUMSAM_COLUMN_WIDTHS):
                assert sheet.column_dimensions[letter].width == width
            assert "적용기준" in workbook.sheetnames
            assert "공종별인부" in workbook.sheetnames
            rules = [workbook["적용기준"].cell(row, 1).value for row in range(1, 40)]
            assert any(value and "품셈 칸은 표준품셈 원표" in str(value) for value in rules)
        finally:
            workbook.close()


def test_keyword_column_uses_square_mm_not_ascii(tmp_path: Path) -> None:
    assert "㎟" in display_keyword("HIV전선", "14 mm2")
    assert " " in display_keyword("HIV전선", "14 mm2")
    assert display_keyword("HIV전선", "14 mm2") == "HIV전선 14 ㎟"
    assert "mm2" not in display_keyword("HIV전선", "14 mm2").lower()
    rows = [row for row in default_pumsam_rows() if "HIV전선" in str(row.get("명칭") or "")][:8]
    path = save_pumsam_database(rows, directory=tmp_path, discipline="전기")
    workbook = load_workbook(path)
    try:
        sheet = workbook["품셈표"]
        keywords = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
        assert keywords
        assert any("㎟" in str(value or "") for value in keywords)
        assert any(str(value or "").startswith("HIV전선 ") for value in keywords)
        assert all("mm2" not in str(value or "").lower() for value in keywords)
        specs = [sheet.cell(row, 3).value for row in range(2, sheet.max_row + 1)]
        assert all("mm2" not in str(value or "").lower() for value in specs)
    finally:
        workbook.close()


def test_bundled_pumsam_keyword_uses_square_mm() -> None:
    from app.paths import bundled_data_dir

    path = bundled_data_dir() / "전기_표준품셈.xlsx"
    if not path.exists():
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        keywords = []
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
            keywords.append(row[0])
            if len(keywords) >= 400:
                break
        assert any("㎟" in str(value or "") for value in keywords)
        assert all("mm2" not in str(value or "").lower() for value in keywords)
    finally:
        workbook.close()


def test_same_spec_extra_labor_blanks_name_cells_but_keeps_rows(tmp_path: Path) -> None:
    from app.pumsam import import_pumsam_file

    rows = [
        {
            "명칭": "관로 청소 및 도통시험",
            "규격": "150 이하",
            "단위": "km",
            "노무명칭": "보통인부",
            "품셈": 9.7,
            "할증%": 100,
            "품셈근거": "전기2-11-1",
        },
        {
            "명칭": "관로 청소 및 도통시험",
            "규격": "150 이하",
            "단위": "km",
            "노무명칭": "특고압케이블전공",
            "품셈": 7.28,
            "할증%": 100,
            "품셈근거": "전기2-11-1",
        },
        {
            "명칭": "관로 청소 및 도통시험",
            "규격": "300 이하",
            "단위": "km",
            "노무명칭": "보통인부",
            "품셈": 1.0,
            "할증%": 100,
            "품셈근거": "전기2-11-1",
        },
    ]
    path = save_pumsam_database(rows, directory=tmp_path, discipline="전기")
    workbook = load_workbook(path)
    try:
        sheet = workbook["품셈표"]
        assert sheet["B2"].value == "관로 청소 및 도통시험"
        assert "150" in str(sheet["C2"].value or "")
        assert sheet["E2"].value == "보통인부"
        assert sheet["A3"].value in (None, "")
        assert sheet["B3"].value in (None, "")
        assert "150" in str(sheet["C3"].value or "")
        assert sheet["D3"].value == "km"
        assert sheet["E3"].value == "특고압케이블전공"
        assert sheet["F3"].value == 7.28
        assert sheet["B4"].value == "관로 청소 및 도통시험"
        assert "300" in str(sheet["C4"].value or "")
        assert sheet["E4"].value == "보통인부"
    finally:
        workbook.close()

    loaded = import_pumsam_file(path)
    by_job = {row.get("노무명칭"): row for row in loaded if "150" in str(row.get("규격") or "")}
    assert by_job["보통인부"]["품셈"] == 9.7
    assert by_job["특고압케이블전공"]["품셈"] == 7.28
    assert "관로" in str(by_job["특고압케이블전공"]["명칭"])
    assert "150" in str(by_job["특고압케이블전공"]["규격"])
